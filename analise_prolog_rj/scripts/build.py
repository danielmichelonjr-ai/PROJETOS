import sys; sys.path.insert(0,'scripts')
from prolog import *
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
tab=load_table('data/prolog.xlsx')
s=pd.read_pickle('data/sim.pkl').sort_values('DATA PAGAMENTO')
D=pd.DataFrame([price(find(tab,c),p,nf,detail=True) for c,p,nf in zip(s['CEP ENTREGA'],s['PESO ITENS'],s['VALOR PRODUTOS'])],index=s.index)
top=s.CIDADE.value_counts().head(20).index.tolist()
labels=list(s.FAIXA.cat.categories)
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78')
B=Font(name=F); BB=Font(name=F,bold=True); BLUE=Font(name=F,color='0000FF'); T=Font(name=F,bold=True,size=14)
thin=Side(style='thin',color='BFBFBF'); BOR=Border(top=thin,bottom=thin,left=thin,right=thin)
MON='R$ #,##0.00;-R$ #,##0.00;-'; PCT='0.0%;-0.0%;-'
wb=Workbook()
def hdr(ws,row,cols,start=1):
    for j,c in enumerate(cols):
        x=ws.cell(row,start+j,c); x.font=H; x.fill=HF; x.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); x.border=BOR
def fmt(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name!=F: c.font=Font(name=F,bold=c.font.b if c.font else False,color=c.font.color if c.font else None)
# ---------- Pedidos
wp=wb.active; wp.title='Pedidos'
cols=['Pedido HUB','Data pagamento','Canal','CEP','Cidade','Faixa PROLOG (kg)','Peso (kg)','Valor NF (R$)','Transportadora cotada','Custo atual (R$)',
 'PROLOG frete peso','PROLOG ad valorem','PROLOG GRIS','PROLOG TAS','PROLOG pedágio','PROLOG CTe (R$/kg)','PROLOG outra taxa (R$/kg)','PROLOG subtotal s/ ICMS','PROLOG total c/ ICMS','Diferença (R$)','Diferença %','PROLOG mais barata?','Prazo PROLOG (du)','Top 20?']
hdr(wp,1,cols); N=len(s)
for i,(ix,r) in enumerate(s.iterrows(),start=2):
    d=D.loc[ix]
    vals=[r['PEDIDO HUB'],r['DATA PAGAMENTO'],r['CANAL'],r['CEP ENTREGA'],r['CIDADE'],str(r['FAIXA']),r['PESO ITENS'],r['VALOR PRODUTOS'],r['TRANSPORTADOR COTAÇÃO'],r['ATUAL'],
          d.frete_peso,round(d.advalorem,4),round(d.gris,4),d.tas,d.pedagio,round(d.cte,4),round(d.outra,4)]
    for j,v in enumerate(vals,1): wp.cell(i,j,v)
    wp.cell(i,18,f'=SUM(K{i}:Q{i})'); wp.cell(i,19,f'=R{i}/(1-Resumo!$C$6)')
    wp.cell(i,20,f'=S{i}-J{i}'); wp.cell(i,21,f'=IFERROR(S{i}/J{i}-1,0)'); wp.cell(i,22,f'=IF(S{i}<J{i},"SIM","NÃO")')
    wp.cell(i,23,r['PRAZO_PROLOG']); wp.cell(i,24,'SIM' if r['CIDADE'] in top else 'NÃO')
    wp.cell(i,2).number_format='dd/mm/yyyy'
    for j in list(range(8,21))+[10]: wp.cell(i,j).number_format=MON
    wp.cell(i,8).number_format=MON; wp.cell(i,21).number_format=PCT; wp.cell(i,7).number_format='0.00'
for j,w in enumerate([12,12,18,10,24,10,9,12,24,12]+[12]*14,1): wp.column_dimensions[L(j)].width=w
wp.freeze_panes='A2'; wp.auto_filter.ref=f'A1:X{N+1}'; wp.row_dimensions[1].height=45
R=lambda c:f'Pedidos!${c}$2:${c}${N+1}'
# ---------- Resumo
ws=wb.create_sheet('Resumo',0)
ws['A1']='Simulação PROLOG FILIAL SP — Pedidos RJ com origem de cotação FILIAL SP (jul–set/2026)'; ws['A1'].font=T
ws['A3']='PREMISSAS'; ws['A3'].font=BB
prem=[('Alíquota ICMS SP→RJ (gross-up, tabela marcada "ICMS já calculado? = N")',0.12,'Interestadual SP→RJ 12%. Intelipost soma o ICMS por dentro: total = subtotal / (1 − alíquota).'),
      ('Ad valorem (% NF)',0.001,'Da tabela: FRETE VALOR (%) = 0,1'),('GRIS (% NF)',0.002,'Da tabela: GRIS (%) = 0,2'),('TAS (R$ por pedido)',1,'Da tabela'),
      ('Pedágio (R$ por fração de 100 kg)',4.5,'Da tabela: 4,50 / faixa 100 kg'),('CTe (R$ por kg)',0.06,'Da tabela: CTE FIXO 0,06 com faixa de 1 kg → cobrado por kg'),
      ('Outra taxa (R$ por kg)',0.45,'Da tabela: OUTRA TAXA FIXO 0,45 com faixa de 1 kg → cobrado por kg'),('Cubagem','Não aplicada','Fator de cubagem = 0 na tabela; usado PESO ITENS'),
      ('Peso > 250 kg','Faixa 250 + excedente/kg','Nenhum pedido no recorte passa de 250 kg')]
ws['A5']='Parâmetro'; ws['C5']='Valor'; ws['D5']='Fonte / observação'; hdr(ws,5,['Parâmetro','','Valor','Fonte / observação'])
for k,(a,v,o) in enumerate(prem,6):
    ws.cell(k,1,a); c=ws.cell(k,3,v); c.font=BLUE; ws.cell(k,4,o)
    if isinstance(v,float): c.number_format='0.0%' if v>=0.01 or v in(0.001,0.002) else MON
ws['C6'].fill=PatternFill('solid',fgColor='FFFF00'); ws['C7'].number_format='0.0%'; ws['C8'].number_format='0.0%'
for c in ('C9','C10','C11','C12'): ws[c].number_format=MON
ws['C6'].comment=Comment('Altere para testar outro cenário de ICMS (ex.: 0 se a PROLOG confirmar que o ICMS já está embutido).','Claude')
ws['A16']='Obs.: as taxas (exceto ICMS) estão calculadas por pedido na aba Pedidos (colunas K–Q); a célula C6 recalcula o total.'; ws['A16'].font=Font(name=F,italic=True,color='595959')
r0=18; ws.cell(r0,1,'VISÃO GERAL').font=BB
hdr(ws,r0+1,['Indicador','','Atual (transportadoras cotadas)','PROLOG FILIAL SP (simulado)','Diferença'])
rows=[('Pedidos',f'=COUNTA({R("A")})',f'=COUNTA({R("A")})',None,'0'),
      ('Custo total de frete',f'=SUM({R("J")})',f'=SUM({R("S")})','pct',MON),
      ('Custo médio por pedido',f'=AVERAGE({R("J")})',f'=AVERAGE({R("S")})','pct',MON),
      ('Valor de NF total',f'=SUM({R("H")})',f'=SUM({R("H")})',None,MON),
      ('Frete / NF (participação no custo)',f'=C21/C23',f'=D21/D23','pp',PCT),
      ('% de pedidos em que PROLOG seria mais barata',None,f'=COUNTIF({R("V")},"SIM")/C20',None,PCT)]
for k,(a,x,y,dif,nf) in enumerate(rows,r0+2):
    ws.cell(k,1,a)
    if x: ws.cell(k,3,x).number_format=nf
    ws.cell(k,4,y).number_format=nf
    if dif=='pct': ws.cell(k,5,f'=D{k}/C{k}-1').number_format=PCT
    if dif=='pp': ws.cell(k,5,f'=D{k}-C{k}').number_format='+0.0%;-0.0%'
r1=r0+9; ws.cell(r1,1,'DECOMPOSIÇÃO DO PREÇO MÉDIO PROLOG (por pedido)').font=BB
hdr(ws,r1+1,['Componente','','R$ médio/pedido','% do total PROLOG'])
comp=[('Frete peso','K'),('Ad valorem','L'),('GRIS','M'),('TAS','N'),('Pedágio','O'),('CTe (R$/kg)','P'),('Outra taxa (R$/kg)','Q')]
for k,(a,c) in enumerate(comp,r1+2):
    ws.cell(k,1,a); ws.cell(k,3,f'=AVERAGE({R(c)})').number_format=MON; ws.cell(k,4,f'=C{k}/$C${r1+10}').number_format=PCT
k=r1+9; ws.cell(k,1,'ICMS 12% (gross-up)'); ws.cell(k,3,f'=AVERAGE({R("S")})-AVERAGE({R("R")})').number_format=MON; ws.cell(k,4,f'=C{k}/$C${r1+10}').number_format=PCT
k=r1+10; ws.cell(k,1,'TOTAL PROLOG').font=BB; ws.cell(k,3,f'=SUM(C{r1+2}:C{r1+9})').number_format=MON; ws.cell(k,4,f'=C{k}/$C${k}').number_format=PCT
k=r1+11; ws.cell(k,1,'Custo médio atual (comparação)'); ws.cell(k,3,f'=C22').number_format=MON
r2=r1+13; ws.cell(r2,1,'SENSIBILIDADE — quanto cada premissa pesa (diferença do custo total PROLOG vs atual)').font=BB
hdr(ws,r2+1,['Cenário','','Custo total PROLOG','vs atual','% pedidos PROLOG mais barata'])
J=R('J'); K=R('K'); Q=R('Q'); P=R('P'); RR=R('R')
sc=[('A) Tabela literal (taxas por kg + ICMS 12%)',f'=SUM({R("S")})',f'=SUMPRODUCT(--({R("S")}<{J}))/C20'),
    ('B) Sem ICMS gross-up',f'=SUM({RR})',f'=SUMPRODUCT(--({RR}<{J}))/C20'),
    ('C) Sem as taxas por kg (CTe e outra taxa) + ICMS',f'=(SUM({RR})-SUM({P})-SUM({Q}))/(1-$C$6)',f'=SUMPRODUCT(--(({RR}-{P}-{Q})/(1-$C$6)<{J}))/C20'),
    ('D) Sem taxas por kg e sem ICMS',f'=SUM({RR})-SUM({P})-SUM({Q})',f'=SUMPRODUCT(--(({RR}-{P}-{Q})<{J}))/C20'),
    ('E) Somente frete peso',f'=SUM({K})',f'=SUMPRODUCT(--({K}<{J}))/C20')]
for k,(a,x,y) in enumerate(sc,r2+2):
    ws.cell(k,1,a); ws.cell(k,3,x).number_format=MON; ws.cell(k,4,f'=C{k}/$C$21-1').number_format='+0.0%;-0.0%'; ws.cell(k,5,y).number_format=PCT
ws.column_dimensions['A'].width=62; ws.column_dimensions['B'].width=2; ws.column_dimensions['C'].width=22; ws.column_dimensions['D'].width=24; ws.column_dimensions['E'].width=22
ws.row_dimensions[5].height=20
# ---------- Top 20 cidades
wt=wb.create_sheet('Top 20 Cidades',1)
wt['A1']='Top 20 cidades do RJ (por nº de pedidos) — custo atual x PROLOG FILIAL SP simulado'; wt['A1'].font=T
cols=['#','Cidade','Pedidos','% dos pedidos RJ','Peso médio (kg)','NF média (R$)','Custo médio atual (R$)','PROLOG médio (R$)','Diferença média (R$)','PROLOG vs atual (%)','Frete/NF atual','Frete/NF PROLOG','% pedidos PROLOG mais barata','PROLOG só frete peso (R$)','Frete peso vs atual (%)','Desconto necessário p/ empatar','Prazo PROLOG (du)','Avaliação']
hdr(wt,3,cols); wt.row_dimensions[3].height=48
for k,c in enumerate(top,4):
    wt.cell(k,1,k-3); wt.cell(k,2,c)
    crit=f'{R("E")},$B{k}'
    f={3:f'=COUNTIFS({crit})',4:f'=C{k}/Resumo!$C$20',5:f'=AVERAGEIFS({R("G")},{crit})',6:f'=AVERAGEIFS({R("H")},{crit})',
       7:f'=AVERAGEIFS({R("J")},{crit})',8:f'=AVERAGEIFS({R("S")},{crit})',9:f'=H{k}-G{k}',10:f'=H{k}/G{k}-1',
       11:f'=SUMIFS({R("J")},{crit})/SUMIFS({R("H")},{crit})',12:f'=SUMIFS({R("S")},{crit})/SUMIFS({R("H")},{crit})',
       13:f'=COUNTIFS({crit},{R("V")},"SIM")/C{k}',14:f'=AVERAGEIFS({R("K")},{crit})',15:f'=N{k}/G{k}-1',16:f'=MAX(0,1-G{k}/H{k})',
       17:f'=AVERAGEIFS({R("W")},{crit})',18:f'=IF(J{k}<=0,"Muito bem",IF(J{k}<=0.32,"Melhor posição relativa",IF(J{k}<=0.45,"Intermediária","Precisa melhorar")))'}
    for j,v in f.items(): wt.cell(k,j,v)
e=len(top)+3; t=e+1
wt.cell(t,2,'TOTAL TOP 20').font=BB
crit=f'{R("X")},"SIM"'
tf={3:f'=SUM(C4:C{e})',4:f'=C{t}/Resumo!$C$20',5:f'=AVERAGEIFS({R("G")},{crit})',6:f'=AVERAGEIFS({R("H")},{crit})',7:f'=AVERAGEIFS({R("J")},{crit})',8:f'=AVERAGEIFS({R("S")},{crit})',
    9:f'=H{t}-G{t}',10:f'=H{t}/G{t}-1',11:f'=SUMIFS({R("J")},{crit})/SUMIFS({R("H")},{crit})',12:f'=SUMIFS({R("S")},{crit})/SUMIFS({R("H")},{crit})',
    13:f'=COUNTIFS({crit},{R("V")},"SIM")/C{t}',14:f'=AVERAGEIFS({R("K")},{crit})',15:f'=N{t}/G{t}-1',16:f'=MAX(0,1-G{t}/H{t})',17:f'=AVERAGEIFS({R("W")},{crit})'}
for j,v in tf.items(): x=wt.cell(t,j,v); x.font=BB
for k in range(4,t+1):
    for j,nf in {3:'0',4:PCT,5:'0.0',6:MON,7:MON,8:MON,9:MON,10:'+0.0%;-0.0%',11:PCT,12:PCT,13:PCT,14:MON,15:'+0.0%;-0.0%',16:PCT,17:'0.0'}.items(): wt.cell(k,j).number_format=nf
    for j in range(1,19): wt.cell(k,j).border=BOR
wt.conditional_formatting.add(f'J4:J{e}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
wt.conditional_formatting.add(f'O4:O{e}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='max',end_color='F8696B'))
wt.cell(t+2,1,'Avaliação (coluna R): "Muito bem" = PROLOG mais barata que o atual; "Melhor posição relativa" = até +32%; "Intermediária" = +32% a +45%; "Precisa melhorar" = acima de +45%. Faixas definidas pela distribuição observada (média RJ ≈ +43%).').font=Font(name=F,italic=True,color='595959')
wt.cell(t+3,1,'"PROLOG só frete peso" = valor da faixa de peso da tabela, sem ad valorem, GRIS, TAS, pedágio, taxas por kg e ICMS — mostra onde a base da tabela é competitiva e o peso das taxas adicionais.').font=Font(name=F,italic=True,color='595959')
for j,w in enumerate([4,24,9,9,9,12,13,13,12,11,10,10,12,13,12,12,9,24],1): wt.column_dimensions[L(j)].width=w
wt.freeze_panes='C4'
# ---------- Faixa x Transportadora (top 20)
wf=wb.create_sheet('Faixa x Transportadora',2)
wf['A1']='Top 20 cidades — por faixa de peso: transportadoras atuais x PROLOG simulado nos MESMOS pedidos'; wf['A1'].font=T
car=[c for c in s[s.CIDADE.isin(top)]['TRANSPORTADOR COTAÇÃO'].value_counts().index]
blocks=[('TODAS',None)]+[(c,c) for c in car]
row=3
for name,cv in blocks:
    wf.cell(row,1,f'Transportadora atual: {name}').font=BB; row+=1
    hdr(wf,row,['Faixa (kg)','Pedidos','Peso médio','Custo médio atual (R$)','PROLOG médio (R$)','PROLOG vs atual (%)','Frete/NF atual','Frete/NF PROLOG','% pedidos PROLOG mais barata','PROLOG só frete peso (R$)','Desconto necessário p/ empatar']); wf.row_dimensions[row].height=45
    st=row+1
    for lab in labels+['TOTAL']:
        row+=1; wf.cell(row,1,lab)
        crit=f'{R("X")},"SIM"'+(f',{R("I")},"{cv}"' if cv else '')+(f',{R("F")},$A{row}' if lab!='TOTAL' else '')
        fm={2:f'=COUNTIFS({crit})',3:f'=IFERROR(AVERAGEIFS({R("G")},{crit}),"")',4:f'=IFERROR(AVERAGEIFS({R("J")},{crit}),"")',5:f'=IFERROR(AVERAGEIFS({R("S")},{crit}),"")',
            6:f'=IFERROR(E{row}/D{row}-1,"")',7:f'=IFERROR(SUMIFS({R("J")},{crit})/SUMIFS({R("H")},{crit}),"")',8:f'=IFERROR(SUMIFS({R("S")},{crit})/SUMIFS({R("H")},{crit}),"")',
            9:f'=IFERROR(COUNTIFS({crit},{R("V")},"SIM")/B{row},"")',10:f'=IFERROR(AVERAGEIFS({R("K")},{crit}),"")',11:f'=IFERROR(MAX(0,1-D{row}/E{row}),"")'}
        for j,v in fm.items():
            x=wf.cell(row,j,v); x.border=BOR; x.number_format={2:'0',3:'0.0',4:MON,5:MON,6:'+0.0%;-0.0%',7:PCT,8:PCT,9:PCT,10:MON,11:PCT}[j]
            if lab=='TOTAL': x.font=BB
        wf.cell(row,1).border=BOR
        if lab=='TOTAL': wf.cell(row,1).font=BB
    wf.conditional_formatting.add(f'F{st}:F{row-1}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
    row+=2
for j,w in enumerate([22,9,9,13,13,12,10,10,12,13,13],1): wf.column_dimensions[L(j)].width=w
# ---------- Cidade x Faixa
wc=wb.create_sheet('Cidade x Faixa',3)
wc['A1']='Top 20 cidades × faixa de peso — PROLOG vs custo atual (%)  |  verde = PROLOG mais competitiva, vermelho = precisa melhorar'; wc['A1'].font=T
def matrix(r,title,kind):
    wc.cell(r,1,title).font=BB; hdr(wc,r+1,['Cidade']+labels+['Total'])
    for i,c in enumerate(top,r+2):
        wc.cell(i,1,c).border=BOR
        for j,lab in enumerate(labels+['Total'],2):
            crit=f'{R("E")},$A{i}'+(f',{R("F")},"{lab}"' if lab!='Total' else '')
            if kind=='n': v=f'=COUNTIFS({crit})'; nf='0;-0;""'
            elif kind=='dif': v=f'=IFERROR(AVERAGEIFS({R("S")},{crit})/AVERAGEIFS({R("J")},{crit})-1,"")'; nf='+0%;-0%'
            elif kind=='fp': v=f'=IFERROR(AVERAGEIFS({R("K")},{crit})/AVERAGEIFS({R("J")},{crit})-1,"")'; nf='+0%;-0%'
            elif kind=='atual': v=f'=IFERROR(AVERAGEIFS({R("J")},{crit}),"")'; nf='#,##0.00'
            else: v=f'=IFERROR(AVERAGEIFS({R("S")},{crit}),"")'; nf='#,##0.00'
            x=wc.cell(i,j,v); x.number_format=nf; x.border=BOR
    end=r+1+len(top); rng=f'B{r+2}:{L(len(labels)+1)}{end}'
    if kind in('dif',):
        wc.conditional_formatting.add(rng,ColorScaleRule(start_type='num',start_value=0.15,start_color='63BE7B',mid_type='num',mid_value=0.40,mid_color='FFEB84',end_type='num',end_value=0.70,end_color='F8696B'))
    if kind=='fp':
        wc.conditional_formatting.add(rng,ColorScaleRule(start_type='num',start_value=-0.25,start_color='63BE7B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='num',end_value=0.25,end_color='F8696B'))
    return end+3
r=3
r=matrix(r,'1) Quantidade de pedidos','n')
r=matrix(r,'2) PROLOG total simulado vs custo atual (%) — preço completo com taxas e ICMS','dif')
r=matrix(r,'3) Somente frete peso PROLOG vs custo atual (%) — onde a BASE da tabela já é competitiva','fp')
r=matrix(r,'4) Custo médio atual (R$)','atual')
r=matrix(r,'5) PROLOG total simulado médio (R$)','prolog')
wc.column_dimensions['A'].width=24
for j in range(2,len(labels)+3): wc.column_dimensions[L(j)].width=9
wc.freeze_panes='B3'
for w in wb.worksheets: fmt(w)
wb.save('analise_prolog_rj/Simulacao_PROLOG_FilialSP_RJ_Jul-Set.xlsx')
