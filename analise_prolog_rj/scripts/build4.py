# Versão para envio à PROLOG: concorrente anonimizado como OUTRO PARCEIRO, sem dados internos de pedido
import sys, os; sys.path.insert(0,os.path.dirname(__file__))
import prolog, tsl2
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L
TSL_FILE, PROLOG_FILE, SIM, OUT = sys.argv[1:5]
P=prolog.load_table(PROLOG_FILE); T=tsl2.load(TSL_FILE)
BANDS=prolog.BANDS
def tband(t,j): return t['faixas'][t['bands'].index(float(BANDS[j]))]   # tabela do parceiro tem faixa 10 kg = faixa 20 kg
def reg(t): return {33.42:'Capital',42.72:'Metropolitana'}.get(t['faixas'][0],'Interior')
s=pd.read_pickle(SIM).sort_values('DATA PAGAMENTO')
labels=list(s.FAIXA.cat.categories)[:-1]
top=s.CIDADE.value_counts().head(20).index.tolist()
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); HF2=PatternFill('solid',fgColor='C55A11'); HF3=PatternFill('solid',fgColor='548235')
BB=Font(name=F,bold=True); BLUE=Font(name=F,color='0000FF'); T14=Font(name=F,bold=True,size=14); NOTE=Font(name=F,italic=True,color='595959')
thin=Side(style='thin',color='BFBFBF'); BOR=Border(top=thin,bottom=thin,left=thin,right=thin)
MON='R$ #,##0.00;-R$ #,##0.00;-'; PCT='0.0%;-0.0%;-'; SPCT='+0.0%;-0.0%;0.0%'; YEL=PatternFill('solid',fgColor='FFFF00')
GRN_RED=lambda lo,hi: ColorScaleRule(start_type='num',start_value=lo,start_color='63BE7B',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='num',end_value=hi,end_color='F8696B')
wb=Workbook()
def hdr(ws,row,cols,start=1,fill=HF,h=None):
    for j,c in enumerate(cols):
        x=ws.cell(row,start+j,c); x.font=H; x.fill=fill; x.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); x.border=BOR
    if h: ws.row_dimensions[row].height=h
def put(ws,r,c,v,nf=None,font=None):
    x=ws.cell(r,c,v); x.border=BOR
    if nf: x.number_format=nf
    if font: x.font=font
    return x
# ================= Resumo / premissas
ws=wb.active; ws.title='Resumo'
ws['A1']='PROLOG FILIAL SP x OUTRO PARCEIRO — pedidos RJ, origem FILIAL SP, jul–set/2026'; ws['A1'].font=T14
ws['A3']='PREMISSAS (células em azul podem ser alteradas — todas as abas recalculam)'; ws['A3'].font=BB
hdr(ws,4,['Parâmetro','PROLOG','OUTRO PARCEIRO','Fonte / observação'])
t0=T[0]
prem=[('ICMS por fora (total = subtotal / (1 − alíquota))',0.12,0.12,'Confirmado: 12% por fora. Ambas as tabelas com "ICMS incluso = N".','0.0%'),
 ('Ad valorem / frete valor (% NF)',0.001,t0['adv'],'Tabelas Intelipost','0.00%'),
 ('GRIS (% NF)',0.002,t0['gris'],'Tabelas Intelipost','0.00%'),
 ('Seguro (% NF)',0,t0['seguro'],'Tabela OUTRO PARCEIRO','0.00%'),
 ('TAS / entrega fixo (R$ por pedido)',1,t0['entrega'],'PROLOG: TAS R$ 1,00',MON),
 ('Pedágio (R$ por fração de 100 kg)',4.5,t0['ped'],'Tabelas Intelipost',MON),
 ('TRT (% sobre o subtotal)',0,t0['trt'],'Nenhuma das duas tabelas cobra TRT','0.0%'),
 ('CTe por kg (R$/kg)',0.06,0,'PROLOG: CTE FIXO 0,06 com faixa de 1 kg',MON),
 ('CTe fixo (R$ por pedido)',0,t0['cte'],'OUTRO PARCEIRO: CTE VALOR FIXO R$ 13,55',MON),
 ('Taxa interna (R$ por kg) — levar o produto até a filial SP',0.45,t0['outra'],'Taxa do embarcador, obrigatória e IGUAL para todos os parceiros (não negociável) — lançada como "outra taxa"',MON),
 ('Cubagem (kg/m³)',0,300,'Sem dimensões na base: simulado por peso real','0')]
for k,(a,p,t,o,nf) in enumerate(prem,5):
    put(ws,k,1,a); put(ws,k,2,p,nf,BLUE); put(ws,k,3,t,nf,BLUE); put(ws,k,4,o)
ws['B5'].fill=YEL; ws['C5'].fill=YEL
PR=dict(icms='Resumo!$B$5',icms_t='Resumo!$C$5',adv_p='Resumo!$B$6',adv_t='Resumo!$C$6',gris_p='Resumo!$B$7',gris_t='Resumo!$C$7',seg_t='Resumo!$C$8',
        tas='Resumo!$B$9',ent='Resumo!$C$9',ped_p='Resumo!$B$10',ped_t='Resumo!$C$10',trt='Resumo!$C$11',cte_p='Resumo!$B$12',cte_t='Resumo!$C$13',
        tx_p='Resumo!$B$14',tx_t='Resumo!$C$14')
# ================= Pedidos
wp=wb.create_sheet('Pedidos')
base=['Data pagamento','CEP','CEP (nº)','Cidade','Região OUTRO PARCEIRO','Faixa (kg)','Peso (kg)','Valor NF (R$)']
tcols=['OUTRO PARCEIRO frete peso','OUTRO PARCEIRO ad valorem','OUTRO PARCEIRO GRIS','OUTRO PARCEIRO seguro','OUTRO PARCEIRO pedágio','OUTRO PARCEIRO entrega','OUTRO PARCEIRO CTe fixo','OUTRO PARCEIRO taxa interna','OUTRO PARCEIRO TRT','OUTRO PARCEIRO subtotal','OUTRO PARCEIRO total c/ ICMS']
pcols=['PROLOG frete peso','PROLOG ad valorem','PROLOG GRIS','PROLOG TAS','PROLOG pedágio','PROLOG CTe','PROLOG taxa interna','PROLOG subtotal','PROLOG total c/ ICMS']
tail=['Dif. PROLOG − OUTRO PARCEIRO (R$)','Dif. %','PROLOG mais barata?','Prazo OUTRO PARCEIRO (du)','Prazo PROLOG (du)','Top 20?','PROLOG PROPOSTA c/ ICMS','Proposta mais barata?']
cols=base+tcols+pcols+tail; C={n:L(i) for i,n in enumerate(cols,1)}; CI={n:i for i,n in enumerate(cols,1)}
hdr(wp,1,cols,h=48)
for n in tcols: wp.cell(1,CI[n]).fill=HF2
for n in pcols: wp.cell(1,CI[n]).fill=HF3
for i,(ix,r) in enumerate(s.iterrows(),start=2):
    cep=r['CEP ENTREGA']; tt=tsl2.find(T,cep); pp=prolog.find(P,cep); w=r['PESO ITENS']; j=next(k for k,b in enumerate(BANDS) if w<=b)
    v=dict(zip(base,[r['DATA PAGAMENTO'],cep,int(cep),r['CIDADE'],reg(tt),str(r['FAIXA']),w,r['VALOR PRODUTOS']]))
    I=lambda n:f'{C[n]}{i}'; W=I('Peso (kg)'); NF=I('Valor NF (R$)')
    v.update({'OUTRO PARCEIRO frete peso':tband(tt,j),'OUTRO PARCEIRO ad valorem':f'={NF}*{PR["adv_t"]}','OUTRO PARCEIRO GRIS':f'={NF}*{PR["gris_t"]}','OUTRO PARCEIRO seguro':f'={NF}*{PR["seg_t"]}',
      'OUTRO PARCEIRO pedágio':f'={PR["ped_t"]}*CEILING({W}/100,1)','OUTRO PARCEIRO entrega':f'={PR["ent"]}','OUTRO PARCEIRO CTe fixo':f'={PR["cte_t"]}','OUTRO PARCEIRO taxa interna':f'={PR["tx_t"]}*CEILING({W},1)',
      'OUTRO PARCEIRO TRT':f'=SUM({I("OUTRO PARCEIRO frete peso")}:{I("OUTRO PARCEIRO taxa interna")})*{PR["trt"]}','OUTRO PARCEIRO subtotal':f'=SUM({I("OUTRO PARCEIRO frete peso")}:{I("OUTRO PARCEIRO TRT")})',
      'OUTRO PARCEIRO total c/ ICMS':f'={I("OUTRO PARCEIRO subtotal")}/(1-{PR["icms_t"]})',
      'PROLOG frete peso':pp['faixas'][j],'PROLOG ad valorem':f'={NF}*{PR["adv_p"]}','PROLOG GRIS':f'={NF}*{PR["gris_p"]}','PROLOG TAS':f'={PR["tas"]}',
      'PROLOG pedágio':f'={PR["ped_p"]}*CEILING({W}/100,1)','PROLOG CTe':f'={PR["cte_p"]}*CEILING({W},1)','PROLOG taxa interna':f'={PR["tx_p"]}*CEILING({W},1)',
      'PROLOG subtotal':f'=SUM({I("PROLOG frete peso")}:{I("PROLOG taxa interna")})','PROLOG total c/ ICMS':f'={I("PROLOG subtotal")}/(1-{PR["icms"]})',
      'Dif. PROLOG − OUTRO PARCEIRO (R$)':f'={I("PROLOG total c/ ICMS")}-{I("OUTRO PARCEIRO total c/ ICMS")}','Dif. %':f'={I("PROLOG total c/ ICMS")}/{I("OUTRO PARCEIRO total c/ ICMS")}-1',
      'PROLOG mais barata?':f'=IF({I("PROLOG total c/ ICMS")}<{I("OUTRO PARCEIRO total c/ ICMS")},"SIM","NÃO")','Prazo OUTRO PARCEIRO (du)':tt['prazo'],'Prazo PROLOG (du)':pp['prazo'],
      'Top 20?':'SIM' if r['CIDADE'] in top else 'NÃO',
      'PROLOG PROPOSTA c/ ICMS':f'=({I("PROLOG frete peso")}*(1-Painel!$C$40)+{I("PROLOG ad valorem")}+{I("PROLOG GRIS")}+Painel!$C$41+Painel!$C$42*CEILING({W}/100,1)+Painel!$C$43*CEILING({W},1)+{I("PROLOG taxa interna")})/(1-{PR["icms"]})',
      'Proposta mais barata?':f'=IF({I("PROLOG PROPOSTA c/ ICMS")}<{I("OUTRO PARCEIRO total c/ ICMS")},"SIM","NÃO")'})
    for n,x in v.items():
        c=wp.cell(i,CI[n],x)
        if n in tcols or n in pcols or n in('Valor NF (R$)','Dif. PROLOG − OUTRO PARCEIRO (R$)','PROLOG PROPOSTA c/ ICMS'): c.number_format=MON
    wp.cell(i,1).number_format='dd/mm/yyyy'; wp.cell(i,CI['Dif. %']).number_format=SPCT
N=len(s); wp.freeze_panes='E2'; wp.auto_filter.ref=f'A1:{L(len(cols))}{N+1}'
for j in range(1,len(cols)+1): wp.column_dimensions[L(j)].width=12
wp.column_dimensions['D'].width=22; wp.column_dimensions['C'].hidden=True
R=lambda n:f'Pedidos!${C[n]}$2:${C[n]}${N+1}'
TT,PT,TF,PF='OUTRO PARCEIRO total c/ ICMS','PROLOG total c/ ICMS','OUTRO PARCEIRO frete peso','PROLOG frete peso'
# ---- visão geral
r0=18; ws.cell(r0,1,'VISÃO GERAL — os mesmos 2.223 pedidos cotados pelas duas tabelas').font=BB
hdr(ws,r0+1,['Indicador','OUTRO PARCEIRO','PROLOG FILIAL SP','Diferença'])
vg=[('Pedidos',f'=COUNTA({R("Cidade")})',f'=COUNTA({R("Cidade")})',None,'0'),
    ('Custo total de frete',f'=SUM({R(TT)})',f'=SUM({R(PT)})','p',MON),
    ('Custo médio por pedido',f'=AVERAGE({R(TT)})',f'=AVERAGE({R(PT)})','p',MON),
    ('Frete / NF',f'=SUM({R(TT)})/SUM({R("Valor NF (R$)")})',f'=SUM({R(PT)})/SUM({R("Valor NF (R$)")})','pp',PCT),
    ('% pedidos em que a PROLOG ganharia a cotação',None,f'=COUNTIF({R("PROLOG mais barata?")},"SIM")/B{r0+2}',None,PCT),
    ('Prazo médio (dias úteis)',f'=AVERAGE({R("Prazo OUTRO PARCEIRO (du)")})',f'=AVERAGE({R("Prazo PROLOG (du)")})','d','0.0')]
for k,(a,x,y,d,nf) in enumerate(vg,r0+2):
    put(ws,k,1,a); put(ws,k,2,x,nf); put(ws,k,3,y,nf)
    put(ws,k,4,{'p':f'=C{k}/B{k}-1','pp':f'=C{k}-B{k}','d':f'=C{k}-B{k}'}.get(d),{'p':SPCT,'pp':'+0.0%;-0.0%','d':'+0.0;-0.0'}.get(d))
r1=r0+12; ws.cell(r1,1,'DE ONDE VEM A DIFERENÇA — média por pedido (R$)').font=BB
hdr(ws,r1+1,['Componente','OUTRO PARCEIRO','PROLOG','PROLOG − OUTRO PARCEIRO','% da diferença'])
comp=[('Frete peso',TF,PF),('Ad valorem','OUTRO PARCEIRO ad valorem','PROLOG ad valorem'),('GRIS','OUTRO PARCEIRO GRIS','PROLOG GRIS'),('Seguro','OUTRO PARCEIRO seguro',None),
      ('TAS / entrega','OUTRO PARCEIRO entrega','PROLOG TAS'),('Pedágio','OUTRO PARCEIRO pedágio','PROLOG pedágio'),('CTe (OUTRO PARCEIRO fixo R$ 13,55 / PROLOG R$ 0,06 por kg)','OUTRO PARCEIRO CTe fixo','PROLOG CTe'),
      ('Taxa interna R$ 0,45/kg','OUTRO PARCEIRO taxa interna','PROLOG taxa interna'),('TRT','OUTRO PARCEIRO TRT',None)]
for k,(a,ct,cp) in enumerate(comp,r1+2):
    put(ws,k,1,a); put(ws,k,2,f'=AVERAGE({R(ct)})' if ct else 0,MON); put(ws,k,3,f'=AVERAGE({R(cp)})' if cp else 0,MON)
    put(ws,k,4,f'=C{k}-B{k}',MON); put(ws,k,5,f'=D{k}/$D${r1+13}',SPCT)
k=r1+11; put(ws,k,1,'ICMS 12% por fora'); put(ws,k,2,f'=AVERAGE({R(TT)})-AVERAGE({R("OUTRO PARCEIRO subtotal")})',MON); put(ws,k,3,f'=AVERAGE({R(PT)})-AVERAGE({R("PROLOG subtotal")})',MON); put(ws,k,4,f'=C{k}-B{k}',MON); put(ws,k,5,f'=D{k}/$D${r1+13}',SPCT)
k=r1+13; put(ws,k,1,'TOTAL',font=BB); put(ws,k,2,f'=SUM(B{r1+2}:B{r1+11})',MON,BB); put(ws,k,3,f'=SUM(C{r1+2}:C{r1+11})',MON,BB); put(ws,k,4,f'=C{k}-B{k}',MON,BB); put(ws,k,5,f'=D{k}/$D${k}',SPCT,BB)
ws.conditional_formatting.add(f'D{r1+2}:D{r1+11}',GRN_RED(-15,45))
for c,w in zip('ABCDE',[60,22,20,70,16]): ws.column_dimensions[c].width=w
# ================= Comparativo das tabelas
wc=wb.create_sheet('Comparativo Tabelas',1)
wc['A1']='Estrutura das tabelas — principais diferenças (OUTRO PARCEIRO = tabela atualizada 17/07/2026)'; wc['A1'].font=T14
hdr(wc,3,['Item','PROLOG FILIAL SP','OUTRO PARCEIRO','Quem leva vantagem','Comentário / oportunidade'],h=30)
itens=[('Regionalização do RJ','93 faixas de CEP, 9 níveis de preço','93 faixas de CEP, 3 níveis: Capital, Metropolitana (10 faixas), Interior (82)','—','OUTRO PARCEIRO tem preço único por região; a PROLOG pode ser agressiva onde tem operação própria'),
 ('Frete peso — Capital até 20 kg','R$ 46,98','R$ 33,42','OUTRO PARCEIRO','+41%'),
 ('Frete peso — Capital 20 a 35 kg','R$ 82,21','R$ 42,11','OUTRO PARCEIRO','+95% — maior gap de valor (faixa mais vendida da capital)'),
 ('Frete peso — Capital 40 a 50 kg','R$ 93,97','R$ 53,55','OUTRO PARCEIRO','+75%'),
 ('Frete peso — Capital 50 a 60 kg','R$ 112,30','R$ 58,60','OUTRO PARCEIRO','+92%'),
 ('Frete peso — Capital 60 a 70 kg','R$ 131,02','R$ 63,21','OUTRO PARCEIRO','+107%'),
 ('Taxa fixa por pedido','TAS R$ 1,00','CTe fixo R$ 13,55','PROLOG','R$ 12,55 a menos por pedido — é a única vantagem relevante da PROLOG em taxas'),
 ('CTe por kg','R$ 0,06 / kg','não cobra','OUTRO PARCEIRO','≈ R$ 3,00 por pedido; trocar por fixo ou zerar'),
 ('Taxa interna do embarcador','R$ 0,45 / kg','R$ 0,45 / kg','Empate','Igual para todos os parceiros, não negociável'),
 ('Ad valorem / Seguro','0,10% NF (ad valorem)','0,21% NF (seguro)','PROLOG','≈ R$ 0,70 a favor'),
 ('GRIS','0,20% NF','0,21% NF','PROLOG','Pequeno'),
 ('TRT','não cobra','não cobra','Empate',''),
 ('Pedágio','R$ 4,50 / 100 kg','R$ 1,05 / 100 kg','OUTRO PARCEIRO','R$ 3,45 a mais por pedido — fácil de negociar'),
 ('Cubagem','não aplica (0)','300 kg/m³','PROLOG','Vantagem para volumosos (não medida: base sem dimensões)'),
 ('Excedente acima de 250 kg','R$ 7,71 a 12,05 / kg','R$ 0,96 a 1,34 / kg','OUTRO PARCEIRO','Sem pedidos acima de 250 kg; ainda assim fora de mercado'),
 ('Prazo — Capital e Metropolitana','5 a 7 dias úteis','2 dias úteis','OUTRO PARCEIRO','O OUTRO PARCEIRO oferece 2 du: a PROLOG perdeu a vantagem de prazo nas regiões de maior volume'),
 ('Prazo — Interior','6 a 11 dias úteis','13 dias úteis','PROLOG','2 a 7 dias mais rápido'),
 ]
for k,row in enumerate(itens,4):
    for j,v in enumerate(row,1):
        x=put(wc,k,j,v); x.alignment=Alignment(wrap_text=True,vertical='top')
    wc.cell(k,4).font=Font(name=F,bold=True,color={'PROLOG':'548235','OUTRO PARCEIRO':'C00000'}.get(row[3],'000000'))
for c,w in zip('ABCDE',[30,26,30,14,64]): wc.column_dimensions[c].width=w
# ================= Top 20 cidades
wt=wb.create_sheet('Top 20 Cidades',2)
wt['A1']='Top 20 cidades RJ — preço final (com taxas e ICMS) PROLOG x OUTRO PARCEIRO nos mesmos pedidos'; wt['A1'].font=T14
cols20=['#','Cidade','Região OUTRO PARCEIRO','Pedidos','% pedidos RJ','Peso médio (kg)','NF média','OUTRO PARCEIRO médio','PROLOG médio','PROLOG vs OUTRO PARCEIRO','Frete/NF OUTRO PARCEIRO','Frete/NF PROLOG','% pedidos PROLOG ganharia','Frete peso PROLOG vs OUTRO PARCEIRO','Prazo OUTRO PARCEIRO','Prazo PROLOG','Redução necessária no preço PROLOG p/ empatar','Avaliação']
hdr(wt,3,cols20,h=60)
NF20={4:'0',5:PCT,6:'0.0',7:MON,8:MON,9:MON,10:SPCT,11:PCT,12:PCT,13:PCT,14:SPCT,15:'0.0',16:'0.0',17:PCT}
for k,c in enumerate(top,4):
    cr=f'{R("Cidade")},$B{k}'
    f={1:k-3,2:c,3:f'=INDEX({R("Região OUTRO PARCEIRO")},MATCH($B{k},{R("Cidade")},0))',4:f'=COUNTIFS({cr})',5:f'=D{k}/Resumo!$B$20',6:f'=AVERAGEIFS({R("Peso (kg)")},{cr})',
       7:f'=AVERAGEIFS({R("Valor NF (R$)")},{cr})',8:f'=AVERAGEIFS({R(TT)},{cr})',9:f'=AVERAGEIFS({R(PT)},{cr})',10:f'=I{k}/H{k}-1',
       11:f'=SUMIFS({R(TT)},{cr})/SUMIFS({R("Valor NF (R$)")},{cr})',12:f'=SUMIFS({R(PT)},{cr})/SUMIFS({R("Valor NF (R$)")},{cr})',
       13:f'=COUNTIFS({cr},{R("PROLOG mais barata?")},"SIM")/D{k}',14:f'=SUMIFS({R(PF)},{cr})/SUMIFS({R(TF)},{cr})-1',
       15:f'=AVERAGEIFS({R("Prazo OUTRO PARCEIRO (du)")},{cr})',16:f'=AVERAGEIFS({R("Prazo PROLOG (du)")},{cr})',17:f'=MAX(0,1-H{k}/I{k})',
       18:f'=IF(J{k}<=0,"Muito bem",IF(J{k}<=0.32,"Melhor posição",IF(J{k}<=0.45,"Intermediária","Precisa melhorar")))'}
    for j,v in f.items(): put(wt,k,j,v,NF20.get(j))
e=len(top)+3; t=e+1; cr=f'{R("Top 20?")},"SIM"'
tf={2:'TOTAL TOP 20',4:f'=SUM(D4:D{e})',5:f'=D{t}/Resumo!$B$20',6:f'=AVERAGEIFS({R("Peso (kg)")},{cr})',7:f'=AVERAGEIFS({R("Valor NF (R$)")},{cr})',8:f'=AVERAGEIFS({R(TT)},{cr})',
    9:f'=AVERAGEIFS({R(PT)},{cr})',10:f'=I{t}/H{t}-1',11:f'=SUMIFS({R(TT)},{cr})/SUMIFS({R("Valor NF (R$)")},{cr})',12:f'=SUMIFS({R(PT)},{cr})/SUMIFS({R("Valor NF (R$)")},{cr})',
    13:f'=COUNTIFS({cr},{R("PROLOG mais barata?")},"SIM")/D{t}',14:f'=SUMIFS({R(PF)},{cr})/SUMIFS({R(TF)},{cr})-1',15:f'=AVERAGEIFS({R("Prazo OUTRO PARCEIRO (du)")},{cr})',
    16:f'=AVERAGEIFS({R("Prazo PROLOG (du)")},{cr})',17:f'=MAX(0,1-H{t}/I{t})'}
for j,v in tf.items(): put(wt,t,j,v,NF20.get(j),BB)
wt.conditional_formatting.add(f'J4:J{e}',ColorScaleRule(start_type='min',start_color='63BE7B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='F8696B'))
wt.conditional_formatting.add(f'N4:N{e}',GRN_RED(-0.3,1.0))
wt.cell(t+2,1,'Avaliação: "Muito bem" = PROLOG mais barata; "Melhor posição" = até +32%; "Intermediária" = +32% a +45%; "Precisa melhorar" = acima de +45%.').font=NOTE
for j,w in enumerate([4,24,14,9,9,9,11,11,11,10,10,10,11,11,8,8,14,18],1): wt.column_dimensions[L(j)].width=w
wt.freeze_panes='D4'
# ================= Faixa x Região
wr=wb.create_sheet('Faixa x Região',3)
wr['A1']='Por faixa de peso e região OUTRO PARCEIRO — preço final PROLOG x OUTRO PARCEIRO (mesmos pedidos, top 20 cidades)'; wr['A1'].font=T14
row=3
for rg in ['TODAS','Capital','Metropolitana','Interior']:
    wr.cell(row,1,f'Região: {rg}').font=BB; row+=1
    hdr(wr,row,['Faixa (kg)','Pedidos','OUTRO PARCEIRO médio','PROLOG médio','PROLOG vs OUTRO PARCEIRO','OUTRO PARCEIRO frete peso','PROLOG frete peso','Frete peso PROLOG vs OUTRO PARCEIRO','% pedidos PROLOG ganharia','Redução necessária'],h=45)
    st=row+1
    for lab in labels+['TOTAL']:
        row+=1; cr=f'{R("Top 20?")},"SIM"'+(f',{R("Região OUTRO PARCEIRO")},"{rg}"' if rg!='TODAS' else '')+(f',{R("Faixa (kg)")},$A{row}' if lab!='TOTAL' else '')
        fm={1:lab,2:f'=COUNTIFS({cr})',3:f'=IFERROR(AVERAGEIFS({R(TT)},{cr}),"")',4:f'=IFERROR(AVERAGEIFS({R(PT)},{cr}),"")',5:f'=IFERROR(D{row}/C{row}-1,"")',
            6:f'=IFERROR(AVERAGEIFS({R(TF)},{cr}),"")',7:f'=IFERROR(AVERAGEIFS({R(PF)},{cr}),"")',8:f'=IFERROR(G{row}/F{row}-1,"")',
            9:f'=IFERROR(COUNTIFS({cr},{R("PROLOG mais barata?")},"SIM")/B{row},"")',10:f'=IFERROR(MAX(0,1-C{row}/D{row}),"")'}
        for j,v in fm.items(): put(wr,row,j,v,{2:'0',3:MON,4:MON,5:SPCT,6:MON,7:MON,8:SPCT,9:PCT,10:PCT}.get(j),BB if lab=='TOTAL' else None)
    wr.conditional_formatting.add(f'E{st}:E{row-1}',GRN_RED(-0.3,0.8)); wr.conditional_formatting.add(f'H{st}:H{row-1}',GRN_RED(-0.3,1.2))
    row+=2
for j,w in enumerate([12,9,12,12,11,12,12,12,12,12],1): wr.column_dimensions[L(j)].width=w
# ================= Frete peso top 20
wf=wb.create_sheet('Frete Peso Top 20',3)
wf['A1']='Frete peso por faixa (valor de tabela, antes de taxas e ICMS) — PROLOG x OUTRO PARCEIRO nas top 20 cidades'; wf['A1'].font=T14
def cityrow(c):
    cep=s.loc[s.CIDADE==c,'CEP ENTREGA'].iloc[0]; return prolog.find(P,cep),tsl2.find(T,cep)
r=4; starts=[]
for bi,(title,fill) in enumerate([('1) PROLOG — frete peso (R$)',HF3),('2) OUTRO PARCEIRO — frete peso (R$)',HF2),('3) PROLOG vs OUTRO PARCEIRO (%)',HF)]):
    wf.cell(r,1,title).font=BB; hdr(wf,r+1,['Cidade','Região OUTRO PARCEIRO']+labels,fill=fill,h=30); starts.append(r+2)
    for i,c in enumerate(top,r+2):
        pp,tt=cityrow(c); put(wf,i,1,c); put(wf,i,2,reg(tt))
        for j in range(len(labels)):
            col=3+j
            if bi==0: put(wf,i,col,pp['faixas'][j],MON,BLUE)
            elif bi==1: put(wf,i,col,tband(tt,j),MON,BLUE)
            else: put(wf,i,col,f'={L(col)}{starts[0]+i-starts[2]}/{L(col)}{starts[1]+i-starts[2]}-1',SPCT)
    if bi==2: wf.conditional_formatting.add(f'C{r+2}:{L(2+len(labels))}{r+1+len(top)}',GRN_RED(-0.3,1.2))
    r+=len(top)+3
wf.column_dimensions['A'].width=24; wf.column_dimensions['B'].width=14
for j in range(3,3+len(labels)): wf.column_dimensions[L(j)].width=10
# ================= Proposta
wq=wb.create_sheet('Proposta Tabela PROLOG',4)
wq['A1']='Frete peso ALVO para a PROLOG empatar com a OUTRO PARCEIRO (mantendo as taxas atuais da PROLOG, inclusive a taxa interna)'; wq['A1'].font=T14
wq['A2']='Alvo = subtotal OUTRO PARCEIRO (frete peso + ad valorem + GRIS + seguro + pedágio + entrega + CTe fixo + taxa interna, × (1 + TRT)) − taxas PROLOG (ad valorem + GRIS + TAS + pedágio + CTe/kg + taxa interna), para um pedido de referência por faixa. O ICMS é igual nas duas e se anula.'; wq['A2'].font=NOTE
nb=len(labels); c0=8
for b,(tit,fill) in enumerate([('PROLOG atual — frete peso',HF3),('OUTRO PARCEIRO — frete peso',HF2),('PROLOG ALVO (empate)',HF),('Redução necessária no frete peso PROLOG',HF)]):
    cs=c0+b*nb; wq.merge_cells(start_row=5,start_column=cs,end_row=5,end_column=cs+nb-1)
    x=wq.cell(5,cs,tit); x.font=H; x.fill=fill; x.alignment=Alignment(horizontal='center')
wq.cell(3,c0+2*nb-1,'Peso de referência (kg) →').font=BB; wq.cell(4,c0+2*nb-1,'NF de referência (R$) →').font=BB
for j,lab in enumerate(labels):
    col=c0+2*nb+j; mid=(([0]+BANDS)[j]+BANDS[j])/2
    put(wq,3,col,f'=IFERROR(AVERAGEIFS({R("Peso (kg)")},{R("Faixa (kg)")},"{lab}"),{mid})','0.0')
    put(wq,4,col,f'=IFERROR(AVERAGEIFS({R("Valor NF (R$)")},{R("Faixa (kg)")},"{lab}"),{L(col)}3*SUM({R("Valor NF (R$)")})/SUM({R("Peso (kg)")}))',MON)
hdr(wq,6,['UF','Cidade (tabela PROLOG)','CEP inicial','CEP final','Região OUTRO PARCEIRO','Prazo PROLOG / OUTRO PARCEIRO','Pedidos no período']+labels*4,h=30)
for i,pp in enumerate(P,7):
    tt=tsl2.find(T,pp['cepf'])
    v=['RJ',pp['cidade'],pp['cepi'],pp['cepf'],reg(tt),f"{pp['prazo']} / {tt['prazo']}",f'=COUNTIFS({R("CEP (nº)")},">="&C{i},{R("CEP (nº)")},"<="&D{i})']
    for j,x in enumerate(v,1): put(wq,i,j,x,'00000000' if j in(3,4) else None)
    for j in range(nb):
        ca,ct,cz,cd=[L(c0+b*nb+j) for b in range(4)]; W=f'{cz}$3'; NFr=f'{cz}$4'
        put(wq,i,c0+j,pp['faixas'][j],MON,BLUE); put(wq,i,c0+nb+j,tband(tt,j),MON,BLUE)
        tsl_sub=f'({ct}{i}+{NFr}*({PR["adv_t"]}+{PR["gris_t"]}+{PR["seg_t"]})+{PR["ped_t"]}*CEILING({W}/100,1)+{PR["ent"]}+{PR["cte_t"]}+{PR["tx_t"]}*CEILING({W},1))*(1+{PR["trt"]})'
        pro_fee=f'({NFr}*({PR["adv_p"]}+{PR["gris_p"]})+{PR["tas"]}+{PR["ped_p"]}*CEILING({W}/100,1)+({PR["cte_p"]}+{PR["tx_p"]})*CEILING({W},1))'
        put(wq,i,c0+2*nb+j,f'={tsl_sub}-{pro_fee}',MON); put(wq,i,c0+3*nb+j,f'=1-{cz}{i}/{ca}{i}',SPCT)
last=6+len(P)
wq.conditional_formatting.add(f'{L(c0+3*nb)}7:{L(c0+4*nb-1)}{last}',ColorScaleRule(start_type='num',start_value=0,start_color='63BE7B',mid_type='num',mid_value=0.3,mid_color='FFEB84',end_type='num',end_value=0.6,end_color='F8696B'))
wq.freeze_panes='H7'; wq.column_dimensions['B'].width=26
for j in [1,3,4,5,6,7]: wq.column_dimensions[L(j)].width=11
for j in range(c0,c0+4*nb): wq.column_dimensions[L(j)].width=9.5
wq.cell(last+2,2,'Redução negativa = a PROLOG já está abaixo do empate nessa faixa/cidade. Faixas com pouco ou nenhum volume (acima de 150 kg) devem ser tratadas caso a caso.').font=NOTE
exec(open(os.path.join(os.path.dirname(__file__),'painel.py')).read())
wb.move_sheet('Pedidos',offset=10)
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name!=F):
                c.font=Font(name=F,bold=c.font.b,italic=c.font.i,color=c.font.color,size=c.font.sz)
wb.save(OUT)
