# Painel executivo (executado dentro do build4: usa wb, R, PR, TT, PT, TF, PF, labels, hdr, put, estilos)
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
OP='OUTRO PARCEIRO'
wd=wb.create_sheet('Painel',0)
wd.sheet_view.showGridLines=False
wd['B1']='COMPETITIVIDADE DA TABELA PROLOG FILIAL SP — RIO DE JANEIRO'; wd['B1'].font=Font(name=F,bold=True,size=16,color='1F4E78')
wd['B2']='Base: 2.223 pedidos reais do RJ (jul–set/2026), saída filial SP, cotados pela tabela PROLOG e pela tabela de OUTRO PARCEIRO. Preços finais com ICMS 12% por fora e todas as taxas.'; wd['B2'].font=NOTE
TILE=PatternFill('solid',fgColor='F2F2F2')
Nn=f'COUNTA({R("Cidade")})'
tiles=[('PREÇO PROLOG vs OUTRO PARCEIRO',f'=SUM({R(PT)})/SUM({R(TT)})-1','+0%;-0%','preço final médio por pedido'),
       ('COTAÇÕES QUE A PROLOG GANHARIA HOJE',f'=COUNTIF({R("PROLOG mais barata?")},"SIM")/{Nn}','0.0%','pedidos em que a PROLOG seria a mais barata'),
       ('FRETE PESO PROLOG vs OUTRO PARCEIRO',f'=SUM({R(PF)})/SUM({R(TF)})-1','+0%;-0%','principal causa da diferença'),
       ('REDUÇÃO NECESSÁRIA NO FRETE PESO',f'=1-(SUM({R(OP+" subtotal")})-(SUM({R("PROLOG subtotal")})-SUM({R(PF)})))/SUM({R(PF)})','0%','para empatar, mantendo as taxas atuais')]
for k,(lab,f,nf,cap) in enumerate(tiles):
    c=2+k*3
    for rr in range(4,8):
        for cc in range(c,c+3): wd.cell(rr,cc).fill=TILE
    wd.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+2); wd.merge_cells(start_row=5,start_column=c,end_row=6,end_column=c+2); wd.merge_cells(start_row=7,start_column=c,end_row=7,end_column=c+2)
    x=wd.cell(4,c,lab); x.font=Font(name=F,bold=True,size=9,color='595959'); x.alignment=Alignment(horizontal='center')
    x=wd.cell(5,c,f); x.font=Font(name=F,bold=True,size=26,color='C00000' if k!=1 else '1F4E78'); x.number_format=nf; x.alignment=Alignment(horizontal='center',vertical='center')
    x=wd.cell(7,c,cap); x.font=Font(name=F,italic=True,size=9,color='595959'); x.alignment=Alignment(horizontal='center')
def title(r,t,sub=None):
    wd.cell(r,2,t).font=Font(name=F,bold=True,size=12,color='1F4E78')
    if sub: wd.cell(r,8,sub).font=NOTE
# ---- Quadro 1
title(9,'QUADRO 1 — Competitividade por região')
hdr(wd,10,['Região','Pedidos','% do volume','Preço médio PROLOG','Preço médio OUTRO PARCEIRO','PROLOG vs OUTRO PARCEIRO','Frete peso PROLOG vs OUTRO PARCEIRO','Redução necessária no frete peso','Prazo PROLOG (du)','Prazo OUTRO PARCEIRO (du)','Situação'],start=2,h=45)
RG='Região '+OP
for i,rg in enumerate(['Capital','Metropolitana','Interior','TOTAL'],11):
    cr=f'{R(RG)},"{rg}"' if rg!='TOTAL' else f'{R(RG)},"<>"'
    f={2:rg,3:f'=COUNTIFS({cr})',4:f'=C{i}/{Nn}',5:f'=AVERAGEIFS({R(PT)},{cr})',6:f'=AVERAGEIFS({R(TT)},{cr})',7:f'=E{i}/F{i}-1',
       8:f'=SUMIFS({R(PF)},{cr})/SUMIFS({R(TF)},{cr})-1',
       9:f'=1-(SUMIFS({R(OP+" subtotal")},{cr})-(SUMIFS({R("PROLOG subtotal")},{cr})-SUMIFS({R(PF)},{cr})))/SUMIFS({R(PF)},{cr})',
       10:f'=AVERAGEIFS({R("Prazo PROLOG (du)")},{cr})',11:f'=AVERAGEIFS({R("Prazo "+OP+" (du)")},{cr})',
       12:f'=IF(G{i}<=0,"COMPETITIVO",IF(G{i}<=0.25,"AJUSTE MODERADO","AJUSTE FORTE"))'}
    for j,v in f.items(): put(wd,i,j,v,{3:'0',4:'0%',5:MON,6:MON,7:'+0%;-0%',8:'+0%;-0%',9:'0%',10:'0.0',11:'0.0'}.get(j),BB if rg=='TOTAL' else None)
wd.conditional_formatting.add('L11:L14',FormulaRule(formula=['L11="AJUSTE FORTE"'],fill=PatternFill('solid',fgColor='F8CBAD'),font=Font(bold=True,color='C00000')))
wd.conditional_formatting.add('L11:L14',FormulaRule(formula=['L11="AJUSTE MODERADO"'],fill=PatternFill('solid',fgColor='FFE699')))
wd.conditional_formatting.add('L11:L14',FormulaRule(formula=['L11="COMPETITIVO"'],fill=PatternFill('solid',fgColor='C6EFCE')))
wd.cell(15,2,'Capital = CEPs 20000-001 a 23799-999 (município do Rio). Metropolitana = Niterói, São Gonçalo, Duque de Caxias, São João de Meriti, Nova Iguaçu, Belford Roxo, Queimados, Nilópolis e Mesquita. Interior = demais.').font=NOTE
# ---- Quadro 2
bl=labels[:10]
title(17,'QUADRO 2 — PROLOG vs OUTRO PARCEIRO por faixa de peso (preço final)','verde = PROLOG competitiva · vermelho = precisa reduzir')
hdr(wd,18,['Região']+[f'{b} kg' for b in bl]+['Total'],start=2,h=30)
for i,rg in enumerate(['Capital','Metropolitana','Interior','TODAS'],19):
    put(wd,i,2,rg,font=BB if rg=='TODAS' else None)
    for j,b in enumerate(bl+['Total'],3):
        cr=(f'{R(RG)},"{rg}"' if rg!='TODAS' else f'{R(RG)},"<>"')+(f',{R("Faixa (kg)")},"{b}"' if b!='Total' else '')
        put(wd,i,j,f'=IFERROR(SUMIFS({R(PT)},{cr})/SUMIFS({R(TT)},{cr})-1,"")','+0%;-0%')
put(wd,23,2,'% do volume',font=NOTE)
for j,b in enumerate(bl+['Total'],3): put(wd,23,j,f'=COUNTIFS({R("Faixa (kg)")},"{b}")/{Nn}' if b!='Total' else f'=SUM(C23:{L(j-1)}23)','0%',NOTE)
wd.conditional_formatting.add(f'C19:{L(2+len(bl)+1)}22',ColorScaleRule(start_type='num',start_value=0,start_color='63BE7B',mid_type='num',mid_value=0.3,mid_color='FFEB84',end_type='num',end_value=0.6,end_color='F8696B'))
# ---- Quadro 3
cb=labels[:6]; cities=['Rio de Janeiro','São Gonçalo','Duque de Caxias','Nova Iguaçu','Niterói','Campos dos Goytacazes','Belford Roxo','Macaé']
title(25,'QUADRO 3 — Frete peso da tabela: ATUAL x ALVO para empatar (faixas com 80% do volume)','alvo = valor máximo de frete peso para empatar; para ganhar a cotação, ficar alguns % abaixo')
wd.cell(26,2,'Cidade').font=H; wd.cell(26,2).fill=HF
wd.merge_cells(start_row=26,start_column=2,end_row=27,end_column=2)
for k,b in enumerate(cb):
    c=3+k*2; wd.merge_cells(start_row=26,start_column=c,end_row=26,end_column=c+1)
    x=wd.cell(26,c,f'{b} kg'); x.font=H; x.fill=HF; x.alignment=Alignment(horizontal='center')
    for d,t in enumerate(['Atual','Alvo']):
        x=wd.cell(27,c+d,t); x.font=H; x.fill=HF if d==0 else HF3; x.alignment=Alignment(horizontal='center'); x.border=BOR
nbq=len(labels); c0q=8
for i,city in enumerate(cities,28):
    put(wd,i,2,city)
    for k,b in enumerate(cb):
        j=labels.index(b); m=f"MATCH($B{i},'Proposta Tabela PROLOG'!$B:$B,0)"
        put(wd,i,3+k*2,f"=INDEX('Proposta Tabela PROLOG'!{L(c0q+j)}:{L(c0q+j)},{m})",MON)
        put(wd,i,4+k*2,f"=INDEX('Proposta Tabela PROLOG'!{L(c0q+2*nbq+j)}:{L(c0q+2*nbq+j)},{m})",MON,Font(name=F,bold=True,color='548235'))
wd.cell(36,2,'Tabela completa (93 faixas de CEP × 13 faixas de peso) na aba "Proposta Tabela PROLOG".').font=NOTE
# ---- Quadro 4: simulador
title(38,'QUADRO 4 — SIMULADOR DE PROPOSTA (altere as células amarelas)','o resultado recalcula sobre os 2.223 pedidos')
hdr(wd,39,['Parâmetro da proposta','Valor'],start=2); hdr(wd,39,['Região','Preço médio PROPOSTA','PROPOSTA vs OUTRO PARCEIRO','Cotações ganhas'],start=5,h=30)
inp=[('Redução linear no frete peso (%)',0.33,'0%'),('TAS (R$ por pedido) — hoje R$ 1,00',1,MON),('Pedágio (R$ por 100 kg) — hoje R$ 4,50',1.05,MON),('CTe (R$ por kg) — hoje R$ 0,06',0,'R$ #,##0.00')]
for i,(a,v,nf) in enumerate(inp,40):
    put(wd,i,2,a); x=put(wd,i,3,v,nf,BLUE); x.fill=YEL
for i,rg in enumerate(['Capital','Metropolitana','Interior','TOTAL'],40):
    cr=f'{R(RG)},"{rg}"' if rg!='TOTAL' else f'{R(RG)},"<>"'
    put(wd,i,5,rg,font=BB if rg=='TOTAL' else None); put(wd,i,6,f'=AVERAGEIFS({R("PROLOG PROPOSTA c/ ICMS")},{cr})',MON)
    put(wd,i,7,f'=SUMIFS({R("PROLOG PROPOSTA c/ ICMS")},{cr})/SUMIFS({R(TT)},{cr})-1','+0.0%;-0.0%'); put(wd,i,8,f'=COUNTIFS({cr},{R("Proposta mais barata?")},"SIM")/COUNTIFS({cr})','0%')
wd.conditional_formatting.add('G40:G43',GRN_RED(-0.2,0.5))
wd.cell(44,2,'A taxa interna do embarcador (R$ 0,45/kg) e o ICMS 12% são mantidos — iguais para todos os parceiros.').font=NOTE
title(46,'Cenários prontos de redução do frete peso (com TAS, pedágio e CTe do simulador)')
hdr(wd,47,['Redução no frete peso','PROLOG vs OUTRO PARCEIRO','Cotações ganhas'],start=2)
for i,rd in enumerate([0,0.2,0.25,0.3,0.35,0.4],48):
    put(wd,i,2,rd,'0%',BLUE)
    adj=f'{R(PF)}*($C$40-B{i})/(1-{PR["icms"]})'
    put(wd,i,3,f'=(SUM({R("PROLOG PROPOSTA c/ ICMS")})+SUM({R(PF)})*($C$40-B{i})/(1-{PR["icms"]}))/SUM({R(TT)})-1','+0.0%;-0.0%')
    put(wd,i,4,f'=SUMPRODUCT(--({R("PROLOG PROPOSTA c/ ICMS")}+{adj}<{R(TT)}))/{Nn}','0%')
wd.conditional_formatting.add('C48:C53',GRN_RED(-0.2,0.5))
# ---- Quadro 5: o que fazer
title(55,'QUADRO 5 — O QUE A PROLOG PRECISA FAZER')
acts=[('1','FRETE PESO','É o que mais pesa na diferença. A redução necessária para empatar está no Quadro 1, coluna I. Prioridade: Capital (46% dos pedidos), depois Interior (Macaé e Campos).'),
      ('2','PEDÁGIO','Passar de R$ 4,50 para cerca de R$ 1,05 a cada 100 kg (≈ R$ 3,70 por pedido).'),
      ('3','CTe POR KG','Zerar o CTe de R$ 0,06/kg (≈ R$ 3,00 por pedido).'),
      ('4','MANTER A TAS','A TAS de R$ 1,00 é a maior vantagem da PROLOG: o parceiro cobra CTe fixo de R$ 13,55 por pedido. Não trocar por taxa fixa maior.'),
      ('5','PRAZO','Na Capital e na Metropolitana o parceiro entrega em 2 dias úteis (PROLOG: 5 a 7). No Interior a PROLOG é mais rápida (6 a 11 contra 13).'),
      ('6','EXCEDENTE > 250 kg','Rever o excedente de R$ 7,71 a R$ 12,05/kg (mercado ≈ R$ 1,00 a R$ 1,35/kg).'),
      ('7','FAIXAS LEVES','Até 20 kg a PROLOG já está perto (≈ +10%); na Metropolitana até 20 kg empata. Bom ponto de partida.')]
for i,(n,a,t) in enumerate(acts,56):
    x=put(wd,i,2,f'{n}. {a}'); x.font=Font(name=F,bold=True,color='FFFFFF'); x.fill=HF; x.alignment=Alignment(vertical='center',indent=1)
    wd.merge_cells(start_row=i,start_column=3,end_row=i,end_column=13); x=wd.cell(i,3,t); x.font=Font(name=F); x.alignment=Alignment(wrap_text=True,vertical='center'); x.border=BOR
    wd.row_dimensions[i].height=30
# ---- Gráficos
wr_=wb['Faixa x Região']
ch=BarChart(); ch.type='col'; ch.grouping='clustered'; ch.title='Preço final médio por faixa de peso (top 20 cidades)'; ch.y_axis.title='R$ por pedido'
ch.add_data(Reference(wr_,min_col=3,max_col=4,min_row=4,max_row=14),titles_from_data=True); ch.set_categories(Reference(wr_,min_col=1,min_row=5,max_row=14))
for ser,col in zip(ch.series,['EB6834','2A78D6']): ser.graphicalProperties.solidFill=col; ser.graphicalProperties.line.solidFill=col
ch.height=8; ch.width=17; ch.legend.position='b'; ch.y_axis.majorGridlines.spPr=None
wd.cell(64,2,'GRÁFICOS').font=Font(name=F,bold=True,size=12,color='1F4E78'); wd.add_chart(ch,'B65')
wt_=wb['Top 20 Cidades']
c2=BarChart(); c2.type='bar'; c2.title='PROLOG vs OUTRO PARCEIRO — top 20 cidades (preço final)'
c2.add_data(Reference(wt_,min_col=10,min_row=3,max_row=23),titles_from_data=True); c2.set_categories(Reference(wt_,min_col=2,min_row=4,max_row=23))
c2.series[0].graphicalProperties.solidFill='2A78D6'; c2.series[0].graphicalProperties.line.solidFill='2A78D6'
c2.legend=None; c2.x_axis.scaling.orientation='maxMin'; c2.y_axis.number_format='0%'; c2.height=14; c2.width=15
wd.add_chart(c2,'I65')
wd.column_dimensions['A'].width=2; wd.column_dimensions['B'].width=34
for j in range(3,15): wd.column_dimensions[L(j)].width=13
wd.freeze_panes='A3'
