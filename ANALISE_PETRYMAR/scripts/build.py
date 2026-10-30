import sys, warnings; sys.path.insert(0,'scripts'); warnings.filterwarnings('ignore')
from agg import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
OUT=sys.argv[1]
d.loc[d.CIDADE.eq('')|d.CIDADE.isna(),'CIDADE']='(cidade não informada)'
A=load_table('data/5efb1329-TABELA_ATUAL_PETRYMAR_-_INTELIPOST.xlsx')
N=load_table('data/f9d33a13-TABELA_MULTIMOVEIS_REAJUSTADA_OUT2026_-_PADR_O_INTELIPOST.xlsx')
UFS=['MT','MS','GO','DF','MG']
# ---------- reajuste tabela
m=A.merge(N,on=['CEPI','CEPF'],how='outer',suffixes=('_ATUAL','_NOVA'),indicator=True)
m['UF']=m.UF_NOVA.fillna(m.UF_ATUAL); m['CIDADE']=m.CIDADE_NOVA.fillna(m.CIDADE_ATUAL)
m['Situação']=m._merge.map({'both':'Mantida','right_only':'NOVA faixa de CEP (só na tabela nova)','left_only':'Removida'})
rc=m[['UF','CIDADE','CEPI','CEPF','Situação','PRAZO_ATUAL','PRAZO_NOVA']].copy()
for b in BANDS:
    rc[f'{b}kg ATUAL']=m[f'w{b}_ATUAL']; rc[f'{b}kg NOVA']=m[f'w{b}_NOVA']; rc[f'{b}kg Var%']=m[f'w{b}_NOVA']/m[f'w{b}_ATUAL']-1
rc['Excedente ATUAL']=m.EXC_ATUAL; rc['Excedente NOVA']=m.EXC_NOVA; rc['Excedente Var%']=m.EXC_NOVA/m.EXC_ATUAL-1
rc['Var% média frete-peso']=rc[[f'{b}kg Var%' for b in BANDS]].mean(axis=1)
rc=rc.sort_values(['UF','CIDADE','CEPI'])
b=m[m._merge=='both']
comp=[]
for uf,g in b.groupby('UF'):
    w=pd.concat([g[f'w{x}_NOVA']/g[f'w{x}_ATUAL']-1 for x in BANDS])
    comp.append({'UF':uf,'Faixas de CEP':len(g),
     'Frete-peso Var% mediana':w.median(),'Frete-peso Var% mínima':w.min(),'Frete-peso Var% máxima':w.max(),
     '% faixas CEP com aumento':(g.w50_NOVA>g.w50_ATUAL*1.001).mean(),'% faixas CEP com redução':(g.w50_NOVA<g.w50_ATUAL*0.999).mean(),
     'Ad Valorem ATUAL (%)':g.ADV_ATUAL.median(),'Ad Valorem NOVA (%)':g.ADV_NOVA.median(),
     'GRIS ATUAL (%)':g.GRIS_ATUAL.median(),'GRIS NOVA (%)':g.GRIS_NOVA.median(),
     'Pedágio ATUAL (R$/100kg)':g.PED_ATUAL.median(),'Pedágio NOVA (R$/100kg)':g.PED_NOVA.median(),
     'Taxa Entrega ATUAL (R$)':g.ENT_ATUAL.median(),'Taxa Entrega NOVA (R$)':g.ENT_NOVA.median(),
     'TRT ATUAL (%)':g.TRT_ATUAL.median(),'TRT NOVA (%)':g.TRT_NOVA.median(),
     'Prazo médio ATUAL (d.u.)':g.PRAZO_ATUAL.mean(),'Prazo médio NOVA (d.u.)':g.PRAZO_NOVA.mean(),
     'Faixas com prazo reduzido':int((g.PRAZO_NOVA<g.PRAZO_ATUAL).sum())})
comp=pd.DataFrame(comp)
c=d[d.PETRY_NOVA.notna()&d.PETRY_ATUAL.notna()]
eff=c.groupby('UF').apply(lambda g:pd.Series({'Pedidos simulados':len(g),'Petrymar ATUAL total (R$)':g.PETRY_ATUAL.sum(),'Petrymar NOVA total (R$)':g.PETRY_NOVA.sum(),'Reajuste efetivo NOVA vs ATUAL':g.PETRY_NOVA.sum()/g.PETRY_ATUAL.sum()-1})).reset_index()
own=d[(d.TRANSP=='Petrymar Transportes')&d.PETRY_NOVA.notna()]
effown=own.groupby('UF').apply(lambda g:pd.Series({'Pedidos feitos pela Petrymar':len(g),'Custo pago (R$)':g.CUSTO_REAL.sum(),'Mesmos pedidos na NOVA (R$)':g.PETRY_NOVA.sum(),'Impacto %':g.PETRY_NOVA.sum()/g.CUSTO_REAL.sum()-1,'Impacto R$ no trimestre':g.PETRY_NOVA.sum()-g.CUSTO_REAL.sum()})).reset_index()
# cities with biggest changes (by city, UFs alvo)
cityvar=rc[rc.UF.isin(UFS)&(rc['Situação']=='Mantida')].groupby(['UF','CIDADE'])['Var% média frete-peso'].mean().reset_index()
ordcnt=d.assign(CN=d.CIDADE_TAB).groupby(['UF','CN']).size().rename('Pedidos jul-set').reset_index().rename(columns={'CN':'CIDADE'})
cityvar=cityvar.merge(ordcnt,on=['UF','CIDADE'],how='left').fillna({'Pedidos jul-set':0})
red=cityvar[cityvar['Var% média frete-peso']<-0.005].sort_values(['Pedidos jul-set'],ascending=False)
aum=cityvar[cityvar['Var% média frete-peso']>0.005].sort_values(['Pedidos jul-set'],ascending=False)
# ---------- análises
res=[]
for grp in ['MATRIZ','FILIAL SP + FILIAL MG']:
    for uf in UFS:
        s=summ(d[(d.GRUPO==grp)&(d.UF==uf)]); s['Origem']=grp; s['UF']=uf; res.append(s)
    s=summ(d[d.GRUPO==grp]); s['Origem']=grp; s['UF']='TOTAL'; res.append(s)
res=pd.DataFrame(res); res=res[['Origem','UF']+[x for x in res.columns if x not in('Origem','UF')]]
tops={g:top(d[d.GRUPO==g]) for g in ['MATRIZ','FILIAL SP + FILIAL MG']}
def in_top(grp):
    t=tops[grp]; x=d[d.GRUPO==grp]
    return x[x.set_index(['UF','CIDADE']).index.isin(list(zip(t.UF,t.Cidade)))]
bt={g:band_total(in_top(g),['UF']) for g in tops}
bc={g:band(in_top(g),['UF']) for g in tops}
cb={g:band_total(in_top(g),['UF','CIDADE']).rename(columns={'CIDADE':'Cidade'}) for g in tops}
tri=d[d.MICRO_TRI.notna()]
tri_micro=[]
for grp in ['MATRIZ','FILIAL SP + FILIAL MG','AMBAS']:
    x=tri if grp=='AMBAS' else tri[tri.GRUPO==grp]
    for mi,g in x.groupby('MICRO_TRI'):
        s=summ(g); s['Origem']=grp; s['Microrregião']=mi; tri_micro.append(s)
    s=summ(x); s['Origem']=grp; s['Microrregião']='TOTAL TRIÂNGULO/ALTO PARANAÍBA'; tri_micro.append(s)
tri_micro=pd.DataFrame(tri_micro); tri_micro=tri_micro[['Origem','Microrregião']+[x for x in tri_micro.columns if x not in('Origem','Microrregião')]]
tri_city=[]
for grp in ['MATRIZ','FILIAL SP + FILIAL MG']:
    x=tri[tri.GRUPO==grp]
    for (mi,cty),g in x.groupby(['MICRO_TRI','CIDADE']):
        s=summ(g); s['Origem']=grp; s['Microrregião']=mi; s['Cidade']=cty; tri_city.append(s)
tri_city=pd.DataFrame(tri_city); tri_city=tri_city[['Origem','Microrregião','Cidade']+[x for x in tri_city.columns if x not in('Origem','Microrregião','Cidade')]].sort_values(['Origem','Pedidos'],ascending=[True,False])
tri_band=band_total(tri,['GRUPO']).rename(columns={'GRUPO':'Origem'})
tri_bc=band(tri,['GRUPO']).rename(columns={'GRUPO':'Origem'})
# ---------- base
base=d[['PEDIDO HUB','NOTA','GRUPO','ORIG','DATA PAGAMENTO','CANAL','UF','CIDADE','MICRO_TRI','CEP ENTREGA','PESO ITENS','FAIXA','VALOR PRODUTOS','FRETE PEDIDO','TRANSP','CUSTO_REAL','ICMS','PETRY_ATUAL','PETRY_NOVA','PRAZO_ATUAL','PRAZO_NOVA','CIDADE_TAB']].copy()
base['Dif R$ NOVA - real']=base.PETRY_NOVA-base.CUSTO_REAL
base['Dif % NOVA vs real']=base.PETRY_NOVA/base.CUSTO_REAL-1
base['% Frete/Prod real']=base.CUSTO_REAL/base['VALOR PRODUTOS']
base['% Frete/Prod NOVA']=base.PETRY_NOVA/base['VALOR PRODUTOS']
base['Petrymar NOVA mais barata?']=np.where(base.PETRY_NOVA.isna(),'SEM COBERTURA',np.where(base.PETRY_NOVA<base.CUSTO_REAL,'SIM','NÃO'))
base.columns=['Pedido HUB','Nota','Grupo origem','Origem cotação','Data pagamento','Canal','UF','Cidade','Microrregião Triângulo','CEP','Peso (kg)','Faixa (kg)','Valor produtos','Frete cobrado cliente','Transportadora real','Custo real transportadora','ICMS simulado','Petrymar ATUAL (simulado)','Petrymar NOVA (simulado)','Prazo ATUAL','Prazo NOVA','Cidade na tabela Petrymar','Dif R$ NOVA - real','Dif % NOVA vs real','% Frete/Prod real','% Frete/Prod NOVA','Petrymar NOVA mais barata?']
pd.to_pickle(dict(res=res,tops=tops,bt=bt,bc=bc,cb=cb,tri_micro=tri_micro,tri_city=tri_city,tri_band=tri_band,comp=comp,eff=eff,effown=effown,red=red,aum=aum),'results.pkl')
# ---------- writer
wb=Workbook(); wb.remove(wb.active)
HF=Font(bold=True,color='FFFFFF'); HFILL=PatternFill('solid',fgColor='1F4E78')
FILLS={'MUITO BEM':'C6EFCE','BEM (':'E2EFDA','ATENÇÃO':'FFEB9C','PRECISA':'FFC7CE','Sem cob':'D9D9D9'}
thin=Side(style='thin',color='BFBFBF')
def fmt_for(col):
    cl=col.lower()
    if '%' in col or 'reajuste' in cl or 'redução' in cl or 'impacto %' in cl or cl.startswith('icms'): return '0.0%'
    if any(k in cl for k in ['r$','custo','petrymar','médio','ticket','valor','frete','excedente','kg atual','kg nova','pedágio','entrega']) and 'pedidos' not in cl and 'prazo' not in cl: return '#,##0.00'
    if 'peso' in cl: return '0.0'
    return None
def write(ws,df,r0=1,title=None,note=None):
    if title:
        ws.cell(r0,1,title).font=Font(bold=True,size=13,color='1F4E78'); r0+=1
    if note:
        ws.cell(r0,1,note).font=Font(italic=True,color='595959'); r0+=1
    for j,col in enumerate(df.columns,1):
        c=ws.cell(r0,j,col); c.font=HF; c.fill=HFILL; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    ws.row_dimensions[r0].height=45
    fm=[fmt_for(cn) for cn in df.columns]
    for i,row in enumerate(df.itertuples(index=False),r0+1):
        fill=None
        for v in row:
            if isinstance(v,str):
                for k,f in FILLS.items():
                    if v.startswith(k): fill=PatternFill('solid',fgColor=f)
        for j,v in enumerate(row,1):
            if isinstance(v,float) and np.isnan(v): v=None
            if isinstance(v,(np.integer,)): v=int(v)
            if isinstance(v,(np.floating,)): v=float(v)
            c=ws.cell(i,j,v)
            if fm[j-1] and isinstance(v,(int,float)): c.number_format=fm[j-1]
            c.border=Border(bottom=thin)
            if fill and (df.columns[j-1] in ('Avaliação',) or j<=4): c.fill=fill
    for j,col in enumerate(df.columns,1):
        L=get_column_letter(j)
        if ws.column_dimensions[L].width in (None,13) or ws.column_dimensions[L].width<14:
            ws.column_dimensions[L].width=max(12,min(34,len(str(col))*0.6+4, max([len(str(x)) for x in df.iloc[:,j-1].head(200)]+[8])+2))
    return r0+len(df)+3
def sheet(name,blocks,freeze='A3'):
    ws=wb.create_sheet(name); r=1
    for b in blocks: r=write(ws,*b,r0=r) if False else write(ws,b[0],r,*b[1:])
    ws.freeze_panes=None
    return ws
# store blocks; LEIA-ME written last but placed first
blocks={}
blocks['2-Reajuste Resumo']=[(comp,'COMPARATIVO TABELA ATUAL x TABELA NOVA (REAJUSTADA OUT/2026) - PETRYMAR','Componentes por UF (mediana das faixas de CEP). Frete-peso = valores das faixas 10 a 200 kg.'),
  (eff,'REAJUSTE EFETIVO - todos os pedidos jul-set (MATRIZ + FILIAIS) simulados nas duas tabelas','Mede o efeito combinado de frete-peso + ad valorem + GRIS + pedágio + entrega + TRT.'),
  (effown,'IMPACTO NOS PEDIDOS QUE A PETRYMAR JÁ TRANSPORTOU (jul-set)','Custo efetivamente pago x mesmos pedidos recalculados na tabela nova.'),
  (red.rename(columns={'CIDADE':'Cidade','Var% média frete-peso':'Var% frete-peso'}),'CIDADES COM REDUÇÃO DE FRETE-PESO (ordenado por volume de pedidos)',None),
  (aum.rename(columns={'CIDADE':'Cidade','Var% média frete-peso':'Var% frete-peso'}),'CIDADES COM AUMENTO DE FRETE-PESO (ordenado por volume de pedidos)',None)]
blocks['3-Reajuste por Faixa CEP']=[(rc,'TABELA ATUAL x NOVA - LINHA A LINHA (todas as faixas de CEP)',None)]
blocks['4-Resumo UF x Origem']=[(res,'RESUMO POR UF E ORIGEM DE COTAÇÃO - Simulação 100% Petrymar (tabela NOVA)','Custo real = COTAÇÃO TRANSPORTADOR (custo da transportadora que efetivamente atendeu). % Frete/Produto = custo de frete ÷ valor dos produtos. Médias e totais consideram apenas pedidos com CEP coberto pela Petrymar.')]
for g,code in [('MATRIZ','MATRIZ'),('FILIAL SP + FILIAL MG','FILIAIS')]:
    blocks[f'5-Top20 {code}']=[(tops[g],f'TOP 20 CIDADES POR UF (mais pedidos) - ORIGEM {g}','Avaliação: MUITO BEM ≤ -10% | BEM < 0% | ATENÇÃO até +10% | PRECISA MELHORAR > +10% (Petrymar NOVA vs custo real). "Redução necessária p/ empatar" = desconto que a Petrymar precisaria dar para igualar o custo real.')]
    blocks[f'6-Faixa Peso {code}']=[(bt[g],f'FAIXAS DE PESO - TOP 20 CIDADES DE CADA UF - ORIGEM {g}',None),
                                    (bc[g],f'FAIXAS DE PESO x TRANSPORTADORA QUE ATENDEU - TOP 20 CIDADES - ORIGEM {g}','Comparação nos MESMOS pedidos: custo real da transportadora x quanto a Petrymar NOVA cobraria.')]
    blocks[f'7-Cidade x Faixa {code}']=[(cb[g],f'TOP 20 CIDADES x FAIXA DE PESO - ORIGEM {g}',None)]
blocks['8-Triângulo Mineiro']=[(tri_micro,'TRIÂNGULO MINEIRO / ALTO PARANAÍBA - POR MICRORREGIÃO (IBGE)','66 municípios da mesorregião Triângulo Mineiro/Alto Paranaíba. Microrregiões Uberlândia, Uberaba, Ituiutaba e Frutal = Triângulo; Araxá, Patrocínio e Patos de Minas = Alto Paranaíba.'),
  (tri_city,'TRIÂNGULO MINEIRO - POR CIDADE',None),(tri_band,'TRIÂNGULO MINEIRO - POR FAIXA DE PESO',None),(tri_bc,'TRIÂNGULO MINEIRO - FAIXA DE PESO x TRANSPORTADORA',None)]
blocks['9-Base Simulada']=[(base,'BASE PEDIDO A PEDIDO (MT, MS, GO, DF, MG - MATRIZ, FILIAL SP e FILIAL MG)',None)]
import json; json.dump({},open('/dev/null','w'))
ws0=wb.create_sheet('1-LEIA-ME')
for name,bl in blocks.items():
    ws=wb.create_sheet(name); r=1
    for blk in bl: r=write(ws,blk[0],r,blk[1],blk[2])
    if name=='9-Base Simulada': ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{get_column_letter(base.shape[1])}{base.shape[0]+2}'
    if name=='3-Reajuste por Faixa CEP': ws.freeze_panes='F3'
from leiame import TXT
for i,(t,k) in enumerate(TXT,1):
    c=ws0.cell(i,1,t)
    if k=='T': c.font=Font(bold=True,size=14,color='1F4E78')
    elif k=='H': c.font=Font(bold=True,size=12,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
    c.alignment=Alignment(wrap_text=True,vertical='top')
ws0.column_dimensions['A'].width=180
import pickle; pickle.dump(None,open('/dev/null','wb'))
wb._sheets.sort(key=lambda w:w.title)
wb.save(OUT); print('ok',OUT)
