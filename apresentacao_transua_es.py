import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

SIM, OUT, TAB = sys.argv[1:4]
VER = sys.argv[4] if len(sys.argv) > 4 else 'V1'
SIM_V1 = sys.argv[5] if len(sys.argv) > 5 else None
SIM_CEN = sys.argv[6] if len(sys.argv) > 6 else None
b = pd.read_excel(SIM, sheet_name='Base Simulada')
b = b.rename(columns={'CUSTO ATUAL (COTAÇÃO TRANSP.)': 'ATUAL', 'TRANSUA C/ ICMS 7%': 'TRANSUA'})
GV = ['Serra', 'Vila Velha', 'Vitória', 'Cariacica', 'Viana']
b['REGIAO'] = b.CIDADE.isin(GV).map({True: 'Grande Vitória', False: 'Interior'})
ORD = ['0 - 20', '20 - 40', '40 - 60', '60 - 70', '70 - 90', '90 - 120', '120 - 160', '160 - 200', '200 - 1500']
b['FX'] = b.FAIXA.replace({'160 - 200': '160+', '200 - 1500': '160+'})
FX = ['0 - 20', '20 - 40', '40 - 60', '60 - 70', '70 - 90', '90 - 120', '120 - 160', '160+']
top = b.CIDADE.value_counts().head(20).index.tolist()
t20 = b[b.CIDADE.isin(top)]

NAVY, BLUE, LIGHT, GREY = '1F3864', '2E75B6', 'DDEBF7', 'F2F2F2'
GREEN_F, GREEN_T = 'C6EFCE', '006100'
YEL_F, YEL_T = 'FFEB9C', '9C5700'
ORA_F, ORA_T = 'FCD5B4', '974706'
RED_F, RED_T = 'FFC7CE', '9C0006'
F = 'Arial'
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
R = '"R$" #,##0.00'
R0 = '"R$" #,##0'
P = '+0.0%;-0.0%;0.0%'
P0 = '0%'

wb = Workbook()


def sheet(name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(66 + i)].width = w
    ws.column_dimensions['A'].width = 2
    last = chr(65 + len(widths))
    ws.merge_cells(f'B1:{last}1')
    ws['B1'] = title
    ws['B1'].font = Font(name=F, size=18, bold=True, color='FFFFFF')
    ws['B1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['B1'].alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 36
    ws.merge_cells(f'B2:{last}2')
    ws['B2'] = subtitle
    ws['B2'].font = Font(name=F, size=10, italic=True, color='595959')
    ws['B2'].alignment = Alignment(indent=1)
    return ws


def section(ws, row, text, ncols):
    ws.cell(row, 2, text).font = Font(name=F, size=12, bold=True, color=NAVY)
    for c in range(2, 2 + ncols):
        ws.cell(row, c).border = Border(bottom=Side(style='medium', color=BLUE))


def header(ws, row, labels):
    for i, l in enumerate(labels):
        c = ws.cell(row, 2 + i, l)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 32


def cellw(ws, row, col, v, fmt=None, bold=False, align='center', fill=None):
    c = ws.cell(row, col, v)
    c.font = Font(name=F, size=10, bold=bold)
    c.alignment = Alignment(horizontal=align, vertical='center')
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    return c


def status_rules(ws, rng):
    for txt, f, t in [('Competitivo', GREEN_F, GREEN_T), ('Ajuste fino', YEL_F, YEL_T),
                      ('Precisa melhorar', ORA_F, ORA_T), ('Crítico', RED_F, RED_T)]:
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{txt}"'],
                                      fill=PatternFill('solid', fgColor=f), font=Font(name=F, bold=True, color=t)))


def diff_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator='lessThanOrEqual', formula=['0'],
                                  font=Font(name=F, bold=True, color=GREEN_T)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0.1'],
                                  font=Font(name=F, bold=True, color=RED_T)))


STATUS = '=IF({d}<=0,"Competitivo",IF({d}<=0.1,"Ajuste fino",IF({d}<=0.2,"Precisa melhorar","Crítico")))'


def legend(ws, row):
    ws.cell(row, 2, 'Legenda de status (diferença TRANSUA x mercado atual):').font = Font(name=F, size=9, bold=True, color='595959')
    for i, (txt, f, t, d) in enumerate([('Competitivo', GREEN_F, GREEN_T, '≤ 0%'), ('Ajuste fino', YEL_F, YEL_T, 'até +10%'),
                                        ('Precisa melhorar', ORA_F, ORA_T, '+10% a +20%'), ('Crítico', RED_F, RED_T, 'acima de +20%')]):
        c = ws.cell(row + 1 + i, 2, f'{txt}  ({d})')
        c.font = Font(name=F, size=9, bold=True, color=t)
        c.fill = PatternFill('solid', fgColor=f)
        c.alignment = Alignment(horizontal='center')


def brl(v):
    return f'R$ {v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def pct(v, sign=True):
    return (f'{v:+.1%}' if sign else f'{v:.0%}').replace('.', ',')


def load(path):
    x = pd.read_excel(path, sheet_name='Base Simulada').rename(columns={'CUSTO ATUAL (COTAÇÃO TRANSP.)': 'ATUAL', 'TRANSUA C/ ICMS 7%': 'TRANSUA'})
    x['REGIAO'] = x.CIDADE.isin(GV).map({True: 'Grande Vitória', False: 'Interior'})
    return x


def dif(x):
    return x.TRANSUA.sum() / x.ATUAL.sum() - 1


def win(x):
    return (x.TRANSUA < x.ATUAL).mean()


tb = pd.read_excel(TAB, header=None).iloc[4:, 2:22].astype(float)
tb.columns = ['CEPI', 'CEPF', 'PRAZO', 'MINIMO', 10, 20, 40, 60, 70, 100, 150, 200, 'EXCED', 'ADVALOR', 'SECCAT', 'GRIS', 'TAS', 'PEDAGIO', 'PED_FRAC', 'ENTREGA']
MIN_CAP, MIN_INT = tb.MINIMO.min(), tb.MINIMO.max()
SECCAT = tb.SECCAT.max()
gua = tb[tb.CEPI == 29200000].iloc[0]

# ---------------- Aba 1: RESUMO ----------------
ws = sheet('Resumo', f'TRANSUA  |  Simulação de Frete – Espírito Santo  |  Proposta {VER}',
           'Pedidos com destino ES pagos de jul a set: preço atual praticado x tabela TRANSUA proposta (valores com ICMS)',
           [26, 16, 16, 16, 16, 16, 16])
kpis = [('Pedidos simulados', len(b), '#,##0'),
        ('Custo atual (mercado)', b.ATUAL.sum(), R0),
        (f'Custo com TRANSUA ({VER})', b.TRANSUA.sum(), R0),
        ('Diferença', dif(b), P),
        ('Pedidos em que a TRANSUA é mais barata', win(b), P0)]
spans = [(2, 2), (3, 4), (5, 6), (7, 7), (8, 8)]
for (lab, v, fmt), (c1, c2) in zip(kpis, spans):
    for r in (4, 5):
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    a = ws.cell(4, c1, lab)
    a.font = Font(name=F, size=9, bold=True, color='595959')
    a.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    v_ = ws.cell(5, c1, v)
    v_.number_format = fmt
    v_.font = Font(name=F, size=20, bold=True, color=(RED_T if v > 0 else GREEN_T) if lab == 'Diferença' else NAVY)
    v_.alignment = Alignment(horizontal='center', vertical='center')
    for r in (4, 5):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = PatternFill('solid', fgColor=LIGHT)
ws.row_dimensions[4].height = 28
ws.row_dimensions[5].height = 36

section(ws, 7, f'Resultado por região – proposta {VER}', 7)
header(ws, 8, ['Região', 'Pedidos', 'Preço médio atual', 'Preço médio TRANSUA', 'Diferença', 'TRANSUA mais barata em', 'Status'])
r = 9
for reg in ['Grande Vitória', 'Interior']:
    g = b[b.REGIAO == reg]
    cellw(ws, r, 2, reg + (' (Serra, Vila Velha, Vitória, Cariacica, Viana)' if reg == 'Grande Vitória' else ' (demais cidades)'), align='left', bold=True)
    cellw(ws, r, 3, len(g), '#,##0')
    cellw(ws, r, 4, round(g.ATUAL.mean(), 2), R)
    cellw(ws, r, 5, round(g.TRANSUA.mean(), 2), R)
    cellw(ws, r, 6, f'=E{r}/D{r}-1', P, bold=True)
    cellw(ws, r, 7, win(g), P0)
    cellw(ws, r, 8, STATUS.format(d=f'F{r}'))
    r += 1
cellw(ws, r, 2, 'Total ES', align='left', bold=True, fill=GREY)
cellw(ws, r, 3, '=SUM(C9:C10)', '#,##0', bold=True, fill=GREY)
cellw(ws, r, 4, '=SUMPRODUCT(C9:C10,D9:D10)/C11', R, bold=True, fill=GREY)
cellw(ws, r, 5, '=SUMPRODUCT(C9:C10,E9:E10)/C11', R, bold=True, fill=GREY)
cellw(ws, r, 6, f'=E{r}/D{r}-1', P, bold=True, fill=GREY)
cellw(ws, r, 7, win(b), P0, bold=True, fill=GREY)
cellw(ws, r, 8, STATUS.format(d=f'F{r}'), bold=True)
ws.column_dimensions['B'].width = 50
diff_rules(ws, 'F9:F11'); status_rules(ws, 'H9:H11')

r = 13
if SIM_V1:
    v1 = load(SIM_V1)
    cen = load(SIM_CEN) if SIM_CEN else None
    section(ws, r, 'Evolução da proposta: diferença vs. mercado (negativo = TRANSUA mais barata)', 7); r += 1
    header(ws, r, ['Região', 'Proposta V1', f'Proposta {VER}', 'Melhora (p.p.)', 'Mais barata em (V1)', f'Mais barata em ({VER})',
                   f'{VER} + ajustes 1 a 3 abaixo' if cen is not None else '']); r += 1
    e0 = r
    for reg in ['Total ES', 'Grande Vitória', 'Interior']:
        f1 = v1 if reg == 'Total ES' else v1[v1.REGIAO == reg]
        f2 = b if reg == 'Total ES' else b[b.REGIAO == reg]
        cellw(ws, r, 2, reg, align='left', bold=True)
        cellw(ws, r, 3, round(dif(f1), 4), P)
        cellw(ws, r, 4, round(dif(f2), 4), P, bold=True)
        cellw(ws, r, 5, f'=(C{r}-D{r})*100', '0.0" p.p."')
        cellw(ws, r, 6, round(win(f1), 3), P0)
        cellw(ws, r, 7, round(win(f2), 3), P0, bold=True)
        if cen is not None:
            f3 = cen if reg == 'Total ES' else cen[cen.REGIAO == reg]
            cellw(ws, r, 8, round(dif(f3), 4), P, bold=True, fill='EBF6EE')
        r += 1
    diff_rules(ws, f'C{e0}:D{r-1}'); diff_rules(ws, f'H{e0}:H{r-1}')
    r += 1

gv = b[b.REGIAO == 'Grande Vitória']; it = b[b.REGIAO == 'Interior']
cid = b.groupby('CIDADE').apply(lambda g: pd.Series({'n': len(g), 'd': dif(g)}), include_groups=False).loc[top]
fxd = t20.groupby('FX').apply(dif, include_groups=False)
comp_fx = [f for f in FX if fxd[f] <= 0.01]
int_ok = cid[(cid.d <= 0.08) & ~cid.index.isin(GV)].index.tolist()
int_bad = cid[(cid.d > 0.15) & ~cid.index.isin(GV)].sort_values('d', ascending=False)
leve_it = it[it['PESO ITENS'] <= 40]
gv20 = gv[gv['PESO ITENS'] <= 20]
sec = b[b.SECCAT > 0]
gu = b[b.CIDADE == 'Guarapari']
pesado = [f for f in FX[5:] if fxd[f] > 0.05]

pos = [f'Grande Vitória: {pct(dif(gv))} vs. mercado; TRANSUA mais barata em {pct(win(gv), False)} dos pedidos. ' +
       ', '.join(f'{c} {pct(cid.d[c])}' for c in GV) + '.',
       f'Faixas de {comp_fx[0].split(" - ")[0]} a {comp_fx[-1].split(" - ")[-1]} kg competitivas no top 20 (' +
       ', '.join(f'{f} kg {pct(fxd[f])}' for f in comp_fx) + ').']
if int_ok:
    pos.append('Interior já próximo do mercado (até +8%): ' + ', '.join(f'{c} {pct(cid.d[c])}' for c in int_ok) + '.')
neg = [f'Interior: {pct(dif(it))} vs. mercado, mais barata em só {pct(win(it), False)} dos pedidos. O frete mínimo de {brl(MIN_INT)} (sem ICMS) '
       f'não mudou: {pct(leve_it.MINIMO_APLICADO.mean(), False)} dos pedidos de até 40 kg do interior caem no mínimo '
       f'(TRANSUA {brl(leve_it.TRANSUA.mean())} x mercado {brl(leve_it.ATUAL.mean())}).',
       f'Guarapari: continua na faixa de CEP do interior (mínimo {brl(gua.MINIMO)}, {gua.PRAZO:.0f} d.u.); {pct(dif(gu))} vs. mercado.',
       'Cidades do interior mais distantes do mercado: ' + ', '.join(f'{c} {pct(v)}' for c, v in int_bad.d.items()) + '.',
       f'Pedidos leves (0–20 kg): {pct(fxd["0 - 20"])}. O mínimo da capital ({brl(MIN_CAP)}) fica acima do mercado na Grande Vitória '
       f'(TRANSUA {brl(gv20.TRANSUA.mean())} x mercado {brl(gv20.ATUAL.mean())}).']
if pesado:
    neg.append('Pedidos pesados: ' + ', '.join(f'{f} kg {pct(fxd[f])}' for f in pesado) + '.')
if len(sec):
    neg.append(f'Taxa SECCAT de {brl(SECCAT)} mantida: {len(sec)} pedidos no período, custo médio TRANSUA {brl(sec.TRANSUA.mean())} x mercado {brl(sec.ATUAL.mean())}.')
asks = [f'Reduzir o frete mínimo do interior de {brl(MIN_INT)} para ~R$ 125,00 (sem ICMS).',
        f'Enquadrar Guarapari (CEPs 29200-000 a 29229-999) na tabela da capital: mínimo {brl(MIN_CAP)} e prazo 12 d.u.',
        f'Retirar a taxa SECCAT de {brl(SECCAT)}.',
        f'Reduzir o mínimo da capital de {brl(MIN_CAP)} para ~R$ 75,00 (sem ICMS) para os pedidos de até 20 kg.']
if pesado:
    asks.append('Reduzir a faixa de 200 kg e o excedente por kg em ~10% a 15% (ver aba "Faixas de Peso").')
if SIM_V1 and cen is not None:
    asks.append(f'Com os ajustes 1 a 3, a proposta {VER} ficaria {pct(dif(cen))} vs. mercado no ES (interior {pct(dif(cen[cen.REGIAO == "Interior"]))}), '
                f'mais barata em {pct(win(cen), False)} dos pedidos.')

section(ws, r, 'Onde a TRANSUA já é competitiva', 7); r += 1
for t in pos:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    c = ws.cell(r, 2, '✔  ' + t); c.font = Font(name=F, size=10, color=GREEN_T); c.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    c.fill = PatternFill('solid', fgColor='EBF6EE'); ws.row_dimensions[r].height = 30; r += 1
r += 1
section(ws, r, 'Onde a TRANSUA ainda precisa melhorar', 7); r += 1
for t in neg:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    c = ws.cell(r, 2, '✖  ' + t); c.font = Font(name=F, size=10, color=RED_T); c.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    c.fill = PatternFill('solid', fgColor='FDEDEF'); ws.row_dimensions[r].height = 30; r += 1
r += 1
section(ws, r, 'Nossa proposta de ajuste', 7); r += 1
for i, t in enumerate(asks, 1):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    last = SIM_V1 and SIM_CEN and i == len(asks)
    c = ws.cell(r, 2, ('➜  ' if last else f'{i}.  ') + t)
    c.font = Font(name=F, size=10, bold=True, color=GREEN_T if last else NAVY); c.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    c.fill = PatternFill('solid', fgColor='EBF6EE' if last else LIGHT); ws.row_dimensions[r].height = 22; r += 1


# ---------------- Aba 2: CIDADES ----------------
ws = sheet('Cidades', 'Top 20 cidades do ES  |  Preço médio por pedido',
           'Cidades com mais pedidos de jul a set. Preço alvo = preço médio praticado hoje pelo mercado nos mesmos pedidos (com ICMS).',
           [26, 11, 10, 16, 16, 12, 14, 16, 18])
header(ws, 4, ['Cidade', 'Região', 'Pedidos', 'Preço alvo (mercado)', 'Preço TRANSUA', 'Diferença', 'TRANSUA mais barata em', 'Redução necessária', 'Status'])
r = 5
for cdd in top:
    g = b[b.CIDADE == cdd]
    cellw(ws, r, 2, cdd, align='left', bold=True)
    cellw(ws, r, 3, 'Capital' if cdd in GV else 'Interior')
    cellw(ws, r, 4, len(g), '#,##0')
    cellw(ws, r, 5, round(g.ATUAL.mean(), 2), R)
    cellw(ws, r, 6, round(g.TRANSUA.mean(), 2), R)
    cellw(ws, r, 7, f'=F{r}/E{r}-1', P, bold=True)
    cellw(ws, r, 8, (g.TRANSUA < g.ATUAL).mean(), P0)
    cellw(ws, r, 9, f'=MAX(0,1-E{r}/F{r})', '0.0%;-0.0%;"—"')
    cellw(ws, r, 10, STATUS.format(d=f'G{r}'))
    if r % 2 == 0:
        for c in range(2, 10):
            ws.cell(r, c).fill = PatternFill('solid', fgColor=GREY)
    r += 1
cellw(ws, r, 2, 'Total top 20', align='left', bold=True, fill=LIGHT)
cellw(ws, r, 3, '', fill=LIGHT)
cellw(ws, r, 4, f'=SUM(D5:D{r-1})', '#,##0', bold=True, fill=LIGHT)
cellw(ws, r, 5, f'=SUMPRODUCT(D5:D{r-1},E5:E{r-1})/D{r}', R, bold=True, fill=LIGHT)
cellw(ws, r, 6, f'=SUMPRODUCT(D5:D{r-1},F5:F{r-1})/D{r}', R, bold=True, fill=LIGHT)
cellw(ws, r, 7, f'=F{r}/E{r}-1', P, bold=True, fill=LIGHT)
cellw(ws, r, 8, (t20.TRANSUA < t20.ATUAL).mean(), P0, bold=True, fill=LIGHT)
cellw(ws, r, 9, f'=MAX(0,1-E{r}/F{r})', '0.0%;-0.0%;"—"', bold=True, fill=LIGHT)
cellw(ws, r, 10, STATUS.format(d=f'G{r}'), bold=True)
diff_rules(ws, f'G5:G{r}'); status_rules(ws, f'J5:J{r}')
ws.conditional_formatting.add(f'I5:I{r-1}', ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                                                           end_type='max', end_color='F8696B'))
legend(ws, r + 2)
ws.freeze_panes = 'C5'

# ---------------- Aba 3: FAIXAS DE PESO ----------------
ws = sheet('Faixas de Peso', 'Faixas de peso  |  Top 20 cidades',
           'Preço médio por pedido em cada faixa de peso (kg). "Melhor preço do mercado" = transportadora mais barata hoje na faixa (média).',
           [24, 10, 16, 16, 16, 12, 14, 16, 18])
header(ws, 4, ['Faixa de peso (kg)', 'Pedidos', 'Preço médio mercado', 'Melhor preço do mercado', 'Preço TRANSUA', 'Diferença x média', 'TRANSUA mais barata em', 'Redução p/ média do mercado', 'Status'])
r = 5
for fx in FX:
    g = t20[t20.FX == fx]
    best = g.groupby('TRANSPORTADOR COTAÇÃO').ATUAL.mean().min()
    cellw(ws, r, 2, fx + ' kg', bold=True)
    cellw(ws, r, 3, len(g), '#,##0')
    cellw(ws, r, 4, round(g.ATUAL.mean(), 2), R)
    cellw(ws, r, 5, round(best, 2), R)
    cellw(ws, r, 6, round(g.TRANSUA.mean(), 2), R)
    cellw(ws, r, 7, f'=F{r}/D{r}-1', P, bold=True)
    cellw(ws, r, 8, (g.TRANSUA < g.ATUAL).mean(), P0)
    cellw(ws, r, 9, f'=MAX(0,1-D{r}/F{r})', '0.0%;-0.0%;"—"')
    cellw(ws, r, 10, STATUS.format(d=f'G{r}'))
    r += 1
fx_end = r - 1
diff_rules(ws, f'G5:G{fx_end}'); status_rules(ws, f'J5:J{fx_end}')
ws.conditional_formatting.add(f'I5:I{fx_end}', ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                                                              end_type='max', end_color='F8696B'))
legend(ws, r + 1)

# gráfico
ch = BarChart(); ch.type = 'col'; ch.style = 10
ch.title = 'Preço médio por faixa de peso: mercado x TRANSUA'
ch.y_axis.title = 'R$ por pedido'; ch.y_axis.numFmt = '"R$" #,##0'; ch.y_axis.majorGridlines = None
data = Reference(ws, min_col=4, max_col=4, min_row=4, max_row=fx_end)
data2 = Reference(ws, min_col=6, max_col=6, min_row=4, max_row=fx_end)
ch.add_data(data, titles_from_data=True); ch.add_data(data2, titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=5, max_row=fx_end))
ch.series[0].graphicalProperties.solidFill = 'A6A6A6'; ch.series[0].graphicalProperties.line.noFill = True
ch.series[1].graphicalProperties.solidFill = BLUE; ch.series[1].graphicalProperties.line.noFill = True
ch.legend.position = 'b'; ch.height = 7.5; ch.width = 22; ch.gapWidth = 60
ch.x_axis.delete = False; ch.y_axis.delete = False
ws.add_chart(ch, f'B{r + 7}')

# mapa de calor cidade x faixa
r = r + 24
section(ws, r, 'Mapa de calor: diferença TRANSUA x mercado por cidade e faixa (verde = TRANSUA mais barata)', 9)
r += 1
header(ws, r, ['Cidade'] + [f + ' kg' for f in FX])
r += 1
h0 = r
for cdd in top:
    cellw(ws, r, 2, cdd, align='left', bold=True)
    for j, fx in enumerate(FX):
        g = b[(b.CIDADE == cdd) & (b.FX == fx)]
        if len(g) >= 3:
            cellw(ws, r, 3 + j, round(g.TRANSUA.sum() / g.ATUAL.sum() - 1, 4), P)
        else:
            cellw(ws, r, 3 + j, '·', fill=GREY).font = Font(name=F, size=9, color='A6A6A6')
    r += 1
ws.conditional_formatting.add(f'C{h0}:J{r-1}', ColorScaleRule(start_type='num', start_value=-0.15, start_color='63BE7B',
                                                             mid_type='num', mid_value=0, mid_color='FFFFFF',
                                                             end_type='num', end_value=0.5, end_color='F8696B'))
n = ws.cell(r, 2, '·  = menos de 3 pedidos na combinação (não exibido).')
n.font = Font(name=F, size=9, italic=True, color='595959')

# ---------------- Aba 4: PREMISSAS ----------------
ws = sheet('Premissas', 'Premissas da simulação', 'Como os valores foram calculados', [110])
lines = [
    ('Base', 'Pedidos com destino ao Espírito Santo, pagos entre 01/07 e 30/09 (1.205 pedidos).'),
    ('Preço atual (mercado)', 'Custo de frete cotado pela transportadora que efetivamente atendeu cada pedido, com ICMS.'),
    ('Preço TRANSUA', 'Tabela proposta (formato Intelipost) aplicada pedido a pedido:'),
    ('', '•  Frete-peso pela faixa (até 10/20/40/60/70/100/150/200 kg; acima de 200 kg = faixa de 200 kg + excedente por kg).'),
    ('', f'•  + Ad valorem {f"{tb.ADVALOR.iat[0]*100:.3f}%".replace(".", ",")} e GRIS {tb.GRIS.iat[0]*100:.1f}% sobre a nota fiscal, TAS {brl(tb.TAS.iat[0])}, pedágio {brl(tb.PEDAGIO.iat[0])} a cada {tb.PED_FRAC.iat[0]:.0f} kg, entrega {brl(tb.ENTREGA.iat[0])} e SECCAT ({brl(SECCAT)} nas faixas de CEP indicadas).'.replace('0.2%', '0,2%')),
    ('', f'•  Frete mínimo: {brl(MIN_CAP)} (Grande Vitória, CEP 29000-000 a 29184-999) e {brl(MIN_INT)} (demais CEPs), aplicado sobre o total.'),
    ('Proposta', f'Tabela simulada: proposta {VER} da TRANSUA' + (' (comparação com a proposta V1 na aba Resumo).' if SIM_V1 else '.')),
    ('ICMS', 'A tabela informa ICMS não incluso; aplicado ICMS interestadual de 7% (origem Sudeste → ES) por dentro: total ÷ 0,93.'),
    ('Peso', 'Peso real dos itens (a base não tem dimensões, portanto a cubagem de 300 kg/m³ não foi aplicada).'),
    ('Confidencialidade', 'Preços de mercado apresentados apenas como médias, sem identificação das transportadoras.'),
]
r = 4
for k, v in lines:
    if k:
        c = ws.cell(r, 2, k); c.font = Font(name=F, size=10, bold=True, color=NAVY)
        r += 1
    c = ws.cell(r, 2, v); c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, indent=1)
    r += 1

del wb['Sheet']
for w in wb.worksheets:
    w.page_setup.orientation = 'landscape'; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb['Resumo'].sheet_properties.tabColor = NAVY
wb.save(OUT)
